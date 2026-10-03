"""05/06 레시피명: 레시피 이름 저장, 06 기준 날짜별 레시피 연결, 검증 오류, 구형 파일 호환."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.db import Base
from app.importer import MigrationImporter
from app.models import MealService, MealServiceMenu, Menu, Recipe
from app.schema_upgrade import upgrade_existing_schema

MENUS = [
    ["M001", "갈비찜", "갈비찜", "주찬", "Y", "정상"],
    ["M002", "배추김치", "배추김치", "김치·절임", "Y", "정상"],
]
INGREDIENTS = [
    ["I001", "갈비", "육류", "kg", 1, "N", "Y", "정상"],
    ["I002", "당근", "채소류", "kg", 1, "N", "Y", "정상"],
    ["I003", "떡볶이떡", "곡류", "kg", 1, "N", "Y", "정상"],
    ["I004", "양파", "채소류", "kg", 1, "N", "Y", "정상"],
    ["I005", "배추김치", "김치·절임", "kg", 1, "N", "Y", "정상"],
]
NAMES = {"M001": "갈비찜", "M002": "배추김치", "I001": "갈비", "I002": "당근", "I003": "떡볶이떡", "I004": "양파", "I005": "배추김치"}

R05_HEADERS = ["메뉴ID", "메뉴명", "레시피명", "재료ID", "표준재료명", "100인기준수량", "단위", "원본행", "원본비고", "검토상태", "재료순서"]
R05_HEADERS_LEGACY = ["메뉴ID", "메뉴명", "재료ID", "표준재료명", "100인기준수량", "단위", "원본행", "원본비고", "검토상태", "재료순서"]
H06 = ["일자", "배식유형", "계획식수", "배식시간", "메뉴순서", "메뉴ID", "메뉴명", "레시피명", "메뉴비고"]
H06_LEGACY = ["일자", "배식유형", "계획식수", "배식시간", "메뉴순서", "메뉴ID", "메뉴명", "메뉴비고"]
H07 = ["일자", "배식유형", "메뉴순서", "메뉴ID", "메뉴명", "재료순서", "재료ID", "표준재료명", "원본재료명", "수량", "단위", "원본비고", "원본행"]

GALBI_A = "갈비찜 - 갈비·당근형"
GALBI_B = "갈비찜 - 갈비·떡볶이떡형"


def recipe_rows(blocks: list[tuple[str, str | None, list[str]]], legacy: bool = False) -> list[list[object]]:
    rows = []
    for menu_code, name, ingredient_codes in blocks:
        for order, code in enumerate(ingredient_codes, start=1):
            row: list[object] = [menu_code, NAMES[menu_code]]
            if not legacy:
                row.append(name or "")
            row += [code, NAMES[code], 1.0, "kg", "", "", "정상", order]
            rows.append(row)
    return rows


def meal_row(day: str, order: int, code: str, recipe_name: str | None = None, legacy: bool = False) -> list[object]:
    row: list[object] = [day, "중식", 400, "12:00", order, code, NAMES[code]]
    if not legacy:
        row.append(recipe_name or "")
    row.append("")
    return row


def daily_rows(day: str, order: int, menu_code: str, ingredient_codes: list[str]) -> list[list[object]]:
    return [
        [day, "중식", order, menu_code, NAMES[menu_code], index, code, NAMES[code], NAMES[code], 4, "kg", "", ""]
        for index, code in enumerate(ingredient_codes, start=1)
    ]


DEFAULT_BLOCKS = [("M001", GALBI_A, ["I001", "I002"]), ("M001", GALBI_B, ["I001", "I003"]), ("M002", "배추김치 - 배추김치형", ["I005"])]


def make_workbook(path: Path, blocks=None, meals=None, daily=None, legacy05: bool = False, legacy06: bool = False) -> Path:
    blocks = DEFAULT_BLOCKS if blocks is None else blocks
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "01_배식설정"
    sheet.append(["배식유형", "기본계획식수", "기본배식시간", "사용여부", "설명"])
    sheet.append(["중식", 400, "12:00", "Y", ""])
    for name, headers, data in [
        ("02_메뉴기준정보", ["메뉴ID", "메뉴명", "통계집계메뉴명", "메뉴역할", "사용여부", "검토상태"], MENUS),
        ("03_재료기준정보", ["재료ID", "표준재료명", "통계분석군", "기본단위", "kg환산계수", "분석제외", "사용여부", "검토상태"], INGREDIENTS),
        ("04_재료별칭_선택", ["원재료별칭", "재료ID", "표준재료명", "출처"], []),
        ("05_메뉴별재료_기준", R05_HEADERS_LEGACY if legacy05 else R05_HEADERS, recipe_rows(blocks, legacy=legacy05)),
        ("06_식단이력_이관", H06_LEGACY if legacy06 else H06, meals or []),
        ("07_식단재료_이관", H07, daily or []),
    ]:
        ws = workbook.create_sheet(name)
        ws.append(headers)
        for row in data:
            ws.append(row)
    workbook.save(path)
    return path


def make_db() -> Session:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    upgrade_existing_schema(engine)
    return Session(engine, expire_on_commit=False)


def recipe_names(db: Session, menu_name: str = "갈비찜") -> list[str]:
    return [
        recipe.name
        for recipe in db.scalars(select(Recipe).join(Menu).where(Menu.name == menu_name).order_by(Recipe.version))
    ]


def linked(db: Session, day: str = "2025-04-01", menu_name: str = "갈비찜") -> MealServiceMenu:
    return db.scalar(
        select(MealServiceMenu)
        .join(MealService)
        .where(MealService.service_date == date.fromisoformat(day), MealServiceMenu.menu_name_snapshot == menu_name)
    )


@pytest.mark.parametrize("mode", ["replace", "merge"])
def test_recipe_names_from_05_are_stored(tmp_path, mode):
    path = make_workbook(tmp_path / "named.xlsx")
    summary, errors = MigrationImporter(path).preview()
    assert errors == [] and summary["recipe_name_column"] is True
    with make_db() as db:
        result = MigrationImporter(path).apply(db, mode=mode)
        assert recipe_names(db) == [GALBI_A, GALBI_B]
        assert recipe_names(db, "배추김치") == ["배추김치 - 배추김치형"]
        assert result["named_recipes"] == 3


def test_06_recipe_name_links_exact_recipe_even_when_07_composition_differs(tmp_path):
    path = make_workbook(
        tmp_path / "link.xlsx",
        meals=[meal_row("2025-04-01", 1, "M001", GALBI_B), meal_row("2025-04-02", 1, "M001", "")],
        daily=[
            # 2025-04-01: 갈비·떡볶이떡 + 양파(비핵심) → composition matches no recipe, 06 says GALBI_B
            *daily_rows("2025-04-01", 1, "M001", ["I001", "I003", "I004"]),
            # 2025-04-02: blank 06 name, 07 exactly equals GALBI_B → composition matching still works
            *daily_rows("2025-04-02", 1, "M001", ["I001", "I003"]),
        ],
    )
    assert MigrationImporter(path).preview()[1] == []
    with make_db() as db:
        result = MigrationImporter(path).apply(db, mode="replace")
        first = linked(db, "2025-04-01")
        assert first.recipe_name_snapshot == GALBI_B
        assert db.get(Recipe, first.recipe_id).name == GALBI_B
        # the day's actual ingredients are kept as recorded (양파 included)
        assert sorted(item.ingredient_name_snapshot for item in first.ingredients) == ["갈비", "떡볶이떡", "양파"]
        assert linked(db, "2025-04-02").recipe_name_snapshot == GALBI_B
        assert result["recipe_links_by_name"] == 1


def test_blank_06_recipe_name_falls_back_to_default_recipe(tmp_path):
    path = make_workbook(
        tmp_path / "blank.xlsx",
        meals=[meal_row("2025-04-01", 1, "M001", "")],
        daily=daily_rows("2025-04-01", 1, "M001", ["I001", "I002", "I004"]),  # matches nothing
    )
    with make_db() as db:
        MigrationImporter(path).apply(db, mode="replace")
        assert linked(db).recipe_name_snapshot == GALBI_A  # first block = default recipe


def test_unknown_06_recipe_name_is_rejected_with_row_number_before_writing(tmp_path):
    path = make_workbook(
        tmp_path / "unknown.xlsx",
        meals=[
            meal_row("2025-04-01", 1, "M001", GALBI_A),
            meal_row("2025-04-01", 2, "M002", GALBI_A),  # name of another menu
            meal_row("2025-04-02", 1, "M001", "갈비찜 - 없는형"),
        ],
    )
    summary, errors = MigrationImporter(path).preview()
    assert summary["ready"] is False
    assert [(error["type"], error["row"]) for error in errors] == [("UNKNOWN_RECIPE_NAME", 3), ("UNKNOWN_RECIPE_NAME", 4)]
    assert "06_식단이력_이관 4행" in errors[1]["message"] and "갈비찜 - 없는형" in errors[1]["message"]
    with make_db() as db:
        db.add(Menu(name="기존 메뉴", canonical_name="기존 메뉴"))
        db.commit()
        with pytest.raises(ValueError, match="05_메뉴별재료_기준의 같은 메뉴ID"):
            MigrationImporter(path).apply(db, mode="replace")
        db.rollback()
        assert db.scalar(select(Menu).where(Menu.name == "기존 메뉴")) is not None


def test_05_name_conflicts_are_rejected(tmp_path):
    same_name_different_set = make_workbook(
        tmp_path / "dup.xlsx",
        blocks=[("M001", GALBI_A, ["I001", "I002"]), ("M001", GALBI_A, ["I001", "I003"])],
    )
    errors = MigrationImporter(same_name_different_set).preview()[1]
    assert [error["type"] for error in errors] == ["DUPLICATE_RECIPE_NAME"]
    assert "2행, 4행" in errors[0]["message"]

    two_names_same_set = make_workbook(
        tmp_path / "same.xlsx",
        blocks=[("M001", GALBI_A, ["I001", "I002"]), ("M001", "갈비찜 - 다른이름형", ["I002", "I001"])],
    )
    errors = MigrationImporter(two_names_same_set).preview()[1]
    assert [error["type"] for error in errors] == ["RECIPE_NAME_SAME_COMPOSITION"]
    with make_db() as db, pytest.raises(ValueError, match="재료 구성이 같은 블록"):
        MigrationImporter(two_names_same_set).apply(db, mode="merge")


def test_same_name_repeated_for_same_composition_is_one_recipe(tmp_path):
    path = make_workbook(
        tmp_path / "repeat.xlsx",
        blocks=[("M001", GALBI_A, ["I001", "I002"]), ("M001", GALBI_B, ["I001", "I003"]), ("M001", GALBI_A, ["I002", "I001"])],
    )
    assert MigrationImporter(path).preview()[1] == []
    with make_db() as db:
        MigrationImporter(path).apply(db, mode="replace")
        assert recipe_names(db) == [GALBI_A, GALBI_B]


def test_consecutive_blocks_split_on_name_change_even_without_order_reset(tmp_path):
    # hand-written 05 where the second block continues 재료순서 (3, 4) instead of restarting at 1
    rows = recipe_rows([("M001", GALBI_A, ["I001", "I002"]), ("M001", GALBI_B, ["I001", "I003"])])
    rows[2][-1], rows[3][-1] = 3, 4
    path = make_workbook(tmp_path / "split.xlsx", blocks=[])
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    for row in rows:
        workbook["05_메뉴별재료_기준"].append(row)
    workbook.save(path)
    with make_db() as db:
        MigrationImporter(path).apply(db, mode="replace")
        assert recipe_names(db) == [GALBI_A, GALBI_B]


def test_legacy_file_without_name_columns_keeps_auto_names_and_composition_linking(tmp_path):
    path = make_workbook(
        tmp_path / "legacy.xlsx",
        meals=[meal_row("2025-04-01", 1, "M001", legacy=True)],
        daily=daily_rows("2025-04-01", 1, "M001", ["I001", "I003"]),
        legacy05=True,
        legacy06=True,
    )
    summary, errors = MigrationImporter(path).preview()
    assert errors == [] and summary["recipe_name_column"] is False and summary["meal_recipe_name_column"] is False
    with make_db() as db:
        result = MigrationImporter(path).apply(db, mode="replace")
        assert recipe_names(db) == ["기본 레시피 v1", "기본 레시피 v2"]
        assert linked(db).recipe_name_snapshot == "기본 레시피 v2"
        assert "named_recipes" not in result and "recipe_links_by_name" not in result


def test_blank_05_name_falls_back_to_auto_name(tmp_path):
    path = make_workbook(tmp_path / "partial.xlsx", blocks=[("M001", GALBI_A, ["I001", "I002"]), ("M001", None, ["I001", "I003"])])
    with make_db() as db:
        MigrationImporter(path).apply(db, mode="replace")
        assert recipe_names(db) == [GALBI_A, "기본 레시피 v2"]


def test_merge_reimport_updates_recipe_names_and_explicit_links(tmp_path):
    legacy = make_workbook(
        tmp_path / "legacy.xlsx",
        meals=[meal_row("2025-04-01", 1, "M001", legacy=True)],
        daily=daily_rows("2025-04-01", 1, "M001", ["I001", "I003", "I004"]),  # falls back to default (v1)
        legacy05=True,
        legacy06=True,
    )
    named = make_workbook(
        tmp_path / "named.xlsx",
        meals=[meal_row("2025-04-01", 1, "M001", GALBI_B)],
        daily=daily_rows("2025-04-01", 1, "M001", ["I001", "I003", "I004"]),
    )
    with make_db() as db:
        MigrationImporter(legacy).apply(db, mode="merge")
        assert recipe_names(db) == ["기본 레시피 v1", "기본 레시피 v2"]
        assert linked(db).recipe_name_snapshot == "기본 레시피 v1"
        recipe_ids = sorted(db.scalars(select(Recipe.id)).all())

        result = MigrationImporter(named).apply(db, mode="merge")
        db.expire_all()
        assert recipe_names(db) == [GALBI_A, GALBI_B]
        assert sorted(db.scalars(select(Recipe.id)).all()) == recipe_ids  # renamed in place, no new recipes
        assert result["recipe_names_updated"] == 3
        item = linked(db)
        assert item.recipe_name_snapshot == GALBI_B
        assert db.get(Recipe, item.recipe_id).name == GALBI_B
