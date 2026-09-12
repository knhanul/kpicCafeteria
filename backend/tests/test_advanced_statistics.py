from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import event, select
from sqlalchemy.orm import Session

from app.advanced_statistics import (
    StationSelectionError,
    data_quality,
    drilldown,
    export_xlsx,
    menu_metrics,
    menu_weather_metrics,
    overview,
    plan_vs_actual,
    station_list,
    weather_metrics,
)
from app.db import Base, create_database_engine
from app.ingredient_statistics import ingredient_statistics
from app.models import Ingredient, MealActual, MealService, MealServiceMenu, MealServiceMenuIngredient, Menu, WeatherHistory
from app.statistics_service import meal_statistics


def make_db(tmp_path):
    engine = create_database_engine(f"sqlite:///{(tmp_path / 'advanced-statistics.db').as_posix()}")
    Base.metadata.create_all(engine)
    return engine


def add_service(db, day, meal_type="LUNCH", planned=100, actual=100, menu=None, representative=True, snapshot=None, note=None):
    service = MealService(service_date=day, meal_type=meal_type, planned_count=planned, note=note)
    db.add(service)
    db.flush()
    if actual != "absent":
        db.add(MealActual(meal_service_id=service.id, actual_count=actual, note=f"actual-{note}" if note else None))
    if menu is not None or snapshot is not None:
        db.add(MealServiceMenu(
            meal_service_id=service.id,
            menu_id=menu.id if menu else None,
            menu_name_snapshot=snapshot or menu.name,
            is_representative=representative,
        ))
    return service


def add_weather(db, day, station="108", name="서울", temp=10, rain=0, humidity=50, snow=0, sunshine=5):
    db.add(WeatherHistory(
        observation_date=day,
        station_id=station,
        station_name=name,
        avg_temp=temp,
        precipitation=rain,
        avg_humidity=humidity,
        snow_depth=snow,
        sunshine_hours=sunshine,
    ))


def test_service_grain_prevents_menu_and_station_duplicate_sums(tmp_path):
    engine = make_db(tmp_path)
    day = date(2025, 1, 1)
    with Session(engine) as db:
        first = Menu(name="A", canonical_name="A")
        second = Menu(name="B", canonical_name="B")
        db.add_all([first, second])
        db.flush()
        service = add_service(db, day, planned=100, actual=80, menu=first)
        db.add(MealServiceMenu(meal_service_id=service.id, menu_id=second.id, menu_name_snapshot="B", is_representative=True))
        add_weather(db, day, "108", "서울")
        add_weather(db, day, "119", "수원")
        db.commit()
        summary = overview(db, day, day)
        selected = weather_metrics(db, day, day, station_id="108")
    assert summary["service_count"] == 1
    assert summary["planned_sum"] == 100
    assert summary["actual_sum"] == 80
    assert selected["matched_service_count"] == 1


def test_exact_plan_actual_formulas_zero_threshold_and_null_exclusion(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 2, 1)
    with Session(engine) as db:
        add_service(db, start, planned=110, actual=100)
        add_service(db, start + timedelta(days=1), planned=50, actual=0)
        add_service(db, start + timedelta(days=2), planned=180, actual=200)
        add_service(db, start + timedelta(days=3), planned=999, actual=None)
        add_service(db, start + timedelta(days=4), planned=999, actual="absent")
        db.commit()
        result = plan_vs_actual(db, start, start + timedelta(days=4))
    assert result["n"] == 3
    assert result["planned_sum"] == 340
    assert result["actual_sum"] == 300
    assert result["mae"] == 26.67
    assert result["wape"] == 26.67
    assert result["bias"] == 13.33
    assert result["bias_rate"] == 13.33
    assert result["over_count"] == 2
    assert result["under_count"] == 1
    assert result["actual_over_plan_count"] == 1
    assert result["actual_under_plan_count"] == 2
    assert result["threshold_n"] == 2
    assert result["within_5_count"] == 0
    assert result["within_5_rate"] == 0
    assert result["within_10_count"] == 2
    assert result["within_10_rate"] == 100


def test_overview_total_plans_series_and_guarded_deterministic_insights(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 1, 6)
    with Session(engine) as db:
        for index in range(6):
            add_service(db, start + timedelta(days=index), "LUNCH", planned=100 + index, actual=90 + index if index < 5 else None)
        add_service(db, start, "DINNER", planned=40, actual=30)
        db.commit()
        result = overview(db, start, start + timedelta(days=5))
        lunch = overview(db, start, start + timedelta(days=4), "lunch")
    assert result["planned_sum"] == sum(range(100, 106)) + 40
    assert result["comparison_planned_sum"] == sum(range(100, 105)) + 40
    assert result["difference"] == result["actual_sum"] - result["comparison_planned_sum"]
    assert len(result["daily"]) == 6
    assert len(result["weekly"]) == 1
    assert len(result["weekday"]) == 7
    assert [item["weekday"] for item in result["weekday"]] == ["월", "화", "수", "목", "금", "토", "일"]
    assert len(result["monthly"]) == 1
    assert [item["key"] for item in result["meal_types"]] == ["LUNCH", "DINNER"]
    assert result["insights"][0]["code"] == "plan_bias"
    assert lunch["insights"][0]["code"] == "plan_bias"


def test_plan_monthly_weekday_distribution_and_iqr_outlier_context(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 3, 1)
    with Session(engine) as db:
        menu = Menu(name="특식", canonical_name="특식")
        db.add(menu)
        db.flush()
        for index in range(10):
            day = start + timedelta(days=index)
            actual = 300 if index == 9 else 100
            add_service(db, day, planned=110, actual=actual, menu=menu if index == 9 else None, note="확인")
            add_weather(db, day, temp=10 + index)
        db.commit()
        result = plan_vs_actual(db, start, start + timedelta(days=9))
    assert result["monthly"][0]["planned_sum"] == 1100
    assert len(result["weekday"]) == 7
    assert sum(item["count"] for item in result["error_distribution"]) == 10
    assert len(result["outliers"]) == 1
    assert result["outliers"][0]["representative_menus"] == ["특식"]
    assert result["outliers"][0]["weather"]["avg_temp"] == 19
    assert result["outliers"][0]["note"] == "actual-확인"


def test_representative_canonical_fallback_dedupe_and_rankings(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 4, 7)
    with Session(engine) as db:
        special = Menu(name="원본", canonical_name="정규메뉴")
        ignored = Menu(name="무시", canonical_name="무시")
        db.add_all([special, ignored])
        db.flush()
        for index in range(5):
            service = add_service(db, start + timedelta(days=index * 7), planned=130, actual=150, menu=special)
            db.add(MealServiceMenu(meal_service_id=service.id, menu_id=special.id, menu_name_snapshot="중복", is_representative=True))
            db.add(MealServiceMenu(meal_service_id=service.id, menu_id=ignored.id, menu_name_snapshot="무시", is_representative=False))
        for index in range(5, 10):
            add_service(db, start + timedelta(days=index * 7), planned=100, actual=100, snapshot="스냅샷메뉴")
        db.commit()
        result = menu_metrics(db, start, start + timedelta(days=63))
    by_name = {item["canonical_name"]: item for item in result["items"]}
    assert set(by_name) == {"정규메뉴", "스냅샷메뉴"}
    assert by_name["정규메뉴"]["occurrence_n"] == 5
    assert by_name["정규메뉴"]["average_planned"] == 130
    assert by_name["정규메뉴"]["average_plan_error"] == -20
    assert by_name["정규메뉴"]["comparator_n"] == 5
    assert by_name["정규메뉴"]["lift_eligible"] is True
    assert by_name["정규메뉴"]["lift"] == 50
    assert result["rankings"]["highest_lift"][0]["canonical_name"] == "정규메뉴"
    assert result["roles"][0]["n"] == 10
    assert result["combinations"]


def test_menu_fallback_to_juchan_role_when_no_representative(tmp_path):
    """When no is_representative=True menus exist, fall back to role='주찬' menus."""
    engine = make_db(tmp_path)
    start = date(2025, 4, 7)
    with Session(engine) as db:
        juchan = Menu(name="불고기", canonical_name="불고기", role="주찬")
        side = Menu(name="김치", canonical_name="김치", role="반찬")
        db.add_all([juchan, side])
        db.flush()
        for index in range(5):
            service = add_service(db, start + timedelta(days=index * 7), planned=100, actual=120)
            # Both menus with is_representative=False (simulating imported data)
            db.add(MealServiceMenu(meal_service_id=service.id, menu_id=juchan.id, menu_name_snapshot="불고기", is_representative=False))
            db.add(MealServiceMenu(meal_service_id=service.id, menu_id=side.id, menu_name_snapshot="김치", is_representative=False))
        for index in range(5, 10):
            add_service(db, start + timedelta(days=index * 7), planned=100, actual=100)
        db.commit()
        result = menu_metrics(db, start, start + timedelta(days=63))
    by_name = {item["canonical_name"]: item for item in result["items"]}
    # Should fall back to 주찬 role: only "불고기" should appear, not "김치"
    assert "불고기" in by_name
    assert "김치" not in by_name
    assert by_name["불고기"]["occurrence_n"] == 5
    assert result["representative_source"] == "주찬_fallback"


def test_menu_lift_ineligible_when_menu_or_comparator_below_five(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 5, 5)
    with Session(engine) as db:
        menu = Menu(name="메뉴", canonical_name="메뉴")
        db.add(menu)
        db.flush()
        for index in range(4):
            add_service(db, start + timedelta(days=index * 7), actual=150, menu=menu)
        for index in range(4, 9):
            add_service(db, start + timedelta(days=index * 7), actual=100)
        db.commit()
        item = menu_metrics(db, start, start + timedelta(days=56))["items"][0]
    assert item["occurrence_n"] == 4
    assert item["comparator_n"] == 5
    assert item["lift_eligible"] is False
    assert item["lift"] is None


def test_exact_weather_bins_null_dry_rain_humidity_snow_scatter_and_pearson(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 6, 1)
    temperatures = [-1, 0, 5, 10, 15, 20, 25, 30, 31, 32, 33, 34]
    with Session(engine) as db:
        for index, temp in enumerate(temperatures):
            day = start + timedelta(days=index)
            add_service(db, day, actual=index * 10)
            if index < 11:
                add_weather(
                    db,
                    day,
                    temp=temp,
                    rain=None if index == 0 else 0 if index == 1 else 1,
                    humidity=5 + index * 10,
                    snow=None if index == 0 else 1 if index == 2 else 0,
                )
        db.commit()
        result = weather_metrics(db, start, start + timedelta(days=11))
    assert [item["bucket"] for item in result["temperature"]] == ["<0", "0~4.9", "5~9.9", "10~14.9", "15~19.9", "20~24.9", "25~29.9", ">=30"]
    assert result["temperature"][0]["average_planned"] == 100
    assert result["temperature"][0]["average_plan_error"] == 100
    rain = {item["bucket"]: item for item in result["rain"]}
    assert rain["NULL"]["n"] == 1
    assert rain["0"]["n"] == 1
    assert rain[">0"]["n"] == 9
    assert result["missing_weather_service_count"] == 1
    assert result["missing_counts"]["precipitation"] == 1
    assert all(item["bucket"] != "missing" for item in result["temperature"])
    assert result["snow"]["meaningful_n"] == 10
    assert len(result["scatter"]) == 11
    assert result["correlations"]["avg_temp"]["n"] == 11
    assert result["correlations"]["avg_temp"]["eligible"] is True
    assert result["correlations"]["avg_temp"]["pearson"] is not None
    assert "N>=10" in result["correlation_disclaimer"]


def test_station_resolution_and_menu_weather_cross_tables(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 7, 7)
    with Session(engine) as db:
        menu = Menu(name="국", canonical_name="국")
        db.add(menu)
        db.flush()
        for index in range(5):
            day = start + timedelta(days=index * 7)
            add_service(db, day, actual=100 + index, menu=menu)
            add_weather(db, day, temp=25, rain=1)
        db.commit()
        stations = station_list(db)
        result = menu_weather_metrics(db, start, start + timedelta(days=28))
    assert stations["auto_selected_station_id"] == "108"
    assert result["weekday_temperature"][0]["weekday"] == "월"
    assert result["weekday_rain"][0]["bucket"] == ">0"
    assert result["representative_menu_temperature"][0]["n"] == 5
    assert result["representative_menu_temperature"][0]["sample_ok"] is True
    assert result["representative_menu_rain"][0]["canonical_name"] == "국"


def test_multiple_stations_require_explicit_selection(tmp_path):
    engine = make_db(tmp_path)
    day = date(2025, 8, 1)
    with Session(engine) as db:
        add_service(db, day, actual=100)
        add_weather(db, day, "108", "서울")
        add_weather(db, day, "119", "수원")
        db.commit()
        try:
            weather_metrics(db, day, day)
        except StationSelectionError:
            pass
        else:
            raise AssertionError("multiple stations must require selection")
        result = weather_metrics(db, day, day, station_id="119")
    assert result["station_name"] == "수원"
    assert result["matched_service_count"] == 1


def test_data_quality_counts_rates_and_canonical_denominator(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 9, 1)
    with Session(engine) as db:
        menu = Menu(name="연결", canonical_name="연결")
        db.add(menu)
        db.flush()
        add_service(db, start, actual=90, menu=menu)
        add_service(db, start + timedelta(days=1), actual=None, snapshot="스냅샷")
        add_service(db, start + timedelta(days=2), actual=0)
        add_weather(db, start, temp=20, rain=0)
        add_weather(db, start + timedelta(days=1), temp=None, rain=None)
        db.commit()
        result = data_quality(db, start, start + timedelta(days=2))
    metrics = result["metrics"]
    assert metrics["actual"] == {"count": 2, "rate": 66.67, "denominator": 3}
    assert metrics["planned"]["rate"] == 100
    assert metrics["representative_menu"]["count"] == 2
    assert metrics["canonical_linkage"] == {"count": 1, "rate": 50.0, "denominator": 2}
    assert metrics["weather_match"]["count"] == 2
    assert metrics["avg_temp"]["count"] == 1
    assert metrics["precipitation"]["count"] == 1


def test_database_pagination_filters_and_summary_detail_export(tmp_path):
    engine = make_db(tmp_path)
    start = date(2025, 10, 1)
    with Session(engine) as db:
        soup = Menu(name="국", canonical_name="국")
        rice = Menu(name="밥", canonical_name="밥")
        db.add_all([soup, rice])
        db.flush()
        add_service(db, start, actual=90, menu=soup)
        add_service(db, start + timedelta(days=1), actual=80, menu=rice)
        add_service(db, start + timedelta(days=2), actual=None, menu=soup)
        add_weather(db, start, temp=25, rain=2)
        add_weather(db, start + timedelta(days=1), temp=5, rain=0)
        add_weather(db, start + timedelta(days=2), temp=27, rain=None)
        db.commit()
        statements = []

        def count_statement(*args):
            statements.append(args[2])

        event.listen(engine, "before_cursor_execute", count_statement)
        page = drilldown(db, start, start + timedelta(days=2), page=1, page_size=1, rain="rain", menu="국", temp_bucket="25~29.9")
        event.remove(engine, "before_cursor_execute", count_statement)
        content = export_xlsx(db, start, start + timedelta(days=2), rain="rain", menu="국", temp_bucket="25~29.9")
    assert page["total"] == 1
    assert len(page["items"]) == 1
    assert len(statements) <= 5
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["요약", "상세자료"]
    assert len(list(workbook["상세자료"].rows)) == 2
    summary = {row[0].value: row[1].value for row in workbook["요약"].rows if row[0].value != "항목"}
    assert summary["서비스수"] == 1
    assert summary["강수필터"] == "rain"


def test_plan_statistics_does_not_choose_between_multiple_weather_stations(tmp_path):
    engine = make_db(tmp_path)
    day = date(2025, 11, 1)
    with Session(engine) as db:
        add_service(db, day, actual=100)
        add_weather(db, day, "108", "서울")
        add_weather(db, day, "119", "수원")
        db.commit()
        result = plan_vs_actual(db, day, day)
    assert result["n"] == 1
    assert result["station_id"] is None
    assert result["outliers"] == []


def test_ingredient_statistics_only_combines_kg_convertible_amounts(tmp_path):
    engine = make_db(tmp_path)
    day = date(2025, 12, 1)
    with Session(engine) as db:
        menu = Menu(name="메뉴", canonical_name="메뉴")
        flour = Ingredient(name="밀가루", stat_group="곡류", default_unit="kg")
        sauce = Ingredient(name="소스", stat_group="소스", default_unit="L")
        excluded = Ingredient(name="제외", stat_group="기타", default_unit="kg", analysis_excluded=True)
        db.add_all([menu, flour, sauce, excluded])
        db.flush()
        service = add_service(db, day, actual=100, menu=menu)
        service_menu = db.scalar(select(MealServiceMenu).where(MealServiceMenu.meal_service_id == service.id))
        db.add_all([
            MealServiceMenuIngredient(meal_service_menu_id=service_menu.id, ingredient_id=flour.id, ingredient_name_snapshot=flour.name, quantity_total=2, unit="kg"),
            MealServiceMenuIngredient(meal_service_menu_id=service_menu.id, ingredient_id=sauce.id, ingredient_name_snapshot=sauce.name, quantity_total=3, unit="L"),
            MealServiceMenuIngredient(meal_service_menu_id=service_menu.id, ingredient_id=excluded.id, ingredient_name_snapshot=excluded.name, quantity_total=10, unit="kg"),
        ])
        db.commit()
        result = ingredient_statistics(db, day, day)
    by_name = {item["ingredient_name"]: item for item in result["top_ingredients"]}
    assert set(by_name) == {"밀가루", "소스"}
    assert by_name["밀가루"]["quantity_kg"] == 2
    assert by_name["소스"]["quantity_kg"] is None
    assert by_name["소스"]["unconverted_amounts"] == {"L": 3.0}


def test_performance_sized_weather_uses_fixed_queries_and_existing_api(tmp_path):
    engine = make_db(tmp_path)
    start = date(2020, 1, 1)
    with Session(engine) as db:
        for index in range(1200):
            day = start + timedelta(days=index)
            add_service(db, day, actual=80 + index % 40)
            add_weather(db, day, temp=index % 30, rain=index % 2)
        db.commit()
        statements = []

        def count_statement(*args):
            statements.append(args[2])

        event.listen(engine, "before_cursor_execute", count_statement)
        result = weather_metrics(db, start, start + timedelta(days=1199))
        event.remove(engine, "before_cursor_execute", count_statement)
        legacy = meal_statistics(db, start, start + timedelta(days=6), "all")
    assert result["matched_service_count"] == 1200
    assert len(statements) <= 3
    assert legacy["summary"]["service_count"] == 7
