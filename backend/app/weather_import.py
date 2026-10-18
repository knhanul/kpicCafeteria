from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.dialects.sqlite import insert as sqlite_insert
from sqlalchemy.orm import Session

from .models import MealPeriodWeather, WeatherHistory, WeatherHourly, WeatherUploadHistory

ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".csv"}
MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_UNCOMPRESSED_SIZE = 100 * 1024 * 1024
MAX_ROWS = 100_000
PREVIEW_TTL = timedelta(minutes=30)
# Hourly observations that make up each meal's serving-time weather (weather_meal_period).
MEAL_WEATHER_HOURS: dict[str, tuple[int, ...]] = {"LUNCH": (11, 12), "DINNER": (17, 18)}
# 1시간 강수량이 확인된 경우에만 사용. 12시 값 = 11~12시, 직전은 10시+11시 값(09~11시).
PRE_MEAL_HOURS: dict[str, tuple[int, ...]] = {"LUNCH": (10, 11), "DINNER": (16, 17)}
# ASOS 시간자료는 일과 시간만 저장한다. 08시와 18시를 포함한다.
ASOS_KEEP_HOURS = range(8, 19)
# 배식 구간 강수는 '그 시각 값이 직전 1시간량'임이 헤더로 확인될 때만 쓴다.
# 기상자료개방포털 ASOS의 '강수량'은 겨울(11~3월)에 3시간 간격이라는 안내만 있고,
# 받은 파일 샘플로 1시간/3시간/일누적을 확정하지 못했다.
ASOS_PRECIP_INTERVAL = "unconfirmed"
INTERPRETATION_VERSION = "2026-10-asos-1"
DAILY_MIN_HOURLY_SAMPLES = 20  # hourly -> daily aggregate only for (nearly) complete days
HOURLY_COMPARE_FIELDS = (
    "station_name", "temperature", "precipitation", "humidity", "wind_speed", "snow_depth",
    "precip_state", "precip_interval", "snow_state", "snow_new_3h", "snow_new_state", "phenomenon", "source_kind",
)
# A later file that simply omits a column must not wipe a value already stored.
_ABSENT_COLUMN_FIELDS = {
    "temperature": ("temperature",),
    "precipitation": ("precipitation", "precip_state", "precip_interval"),
    "humidity": ("humidity",),
    "wind_speed": ("wind_speed",),
    "snow_depth": ("snow_depth", "snow_state"),
    "snow_new_3h": ("snow_new_3h", "snow_new_state"),
    "phenomenon": ("phenomenon",),
}
PRECIP_DAY_NAMES = {"일강수량", "일강수"}
PRECIP_3H_NAMES = {"3시간강수량", "3시간강수"}
PRECIP_1H_NAMES = {"1시간강수량", "1시간강수"}
NULL_MARKERS = {"", "-", "--", "n/a", "na", "null"}
WEATHER_FIELDS = (
    "avg_temp", "min_temp", "max_temp", "precipitation",
    "avg_humidity", "snow_depth", "sunshine_hours",
)
HEADER_ALIASES = {
    "observation_date": {"일시", "날짜", "관측일", "관측일자", "observation_date"},
    "observation_datetime": {"관측일시", "observation_datetime"},
    "station_id": {"지점", "지점번호", "지점코드", "관측지점번호", "관측지점코드", "station_id"},
    "station_name": {"지점명", "관측지점명", "관측소명", "station_name"},
    "avg_temp": {"평균기온", "일평균기온"},
    "min_temp": {"최저기온", "일최저기온"},
    "max_temp": {"최고기온", "일최고기온"},
    "temperature": {"기온", "temperature", "temp"},
    "precipitation": {"일강수량", "강수량", "일강수", "precipitation", "rainfall"},
    "avg_humidity": {"평균습도", "평균상대습도", "일평균상대습도"},
    "humidity": {"습도", "humidity"},
    "snow_depth": {"적설", "최심적설", "일최심적설", "최심신적설", "snow_depth", "snowdepth", "적설snowdepth"},
    "snow_new_3h": {"3시간신적설", "hr3fhsc", "snow_new_3h"},
    "phenomenon": {"현상번호", "일기현상", "phenomenon_code"},
    "temp_qc": {"기온qc플래그", "taqcflag", "ta_qcflg"},
    "precip_qc": {"강수량qc플래그", "rnqcflag", "rn_qcflg"},
    "sunshine_hours": {"일조시간", "합계일조시간", "일조"},
    "wind_speed": {"풍속", "wind_speed"},
    "source_kind": {"자료구분", "source_kind"},
}


class WeatherImportError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower().replace("\n", "").replace("\r", "")
    text = re.sub(r"\([^)]*\)|\[[^]]*\]", "", text)
    text = text.replace("℃", "").replace("°c", "").replace("°", "")
    return re.sub(r"[\s_./-]+", "", text)


NORMALIZED_ALIASES = {
    field: {normalize_header(alias) for alias in aliases}
    for field, aliases in HEADER_ALIASES.items()
}


def map_headers(headers: list[Any]) -> tuple[dict[str, int], list[str]]:
    unknown: list[str] = []
    pending: list[tuple[int, str, str]] = []
    for index, raw in enumerate(headers):
        header = str(raw or "").strip()
        normalized = normalize_header(raw)
        if not normalized:
            continue
        if normalized in PRECIP_3H_NAMES:
            field = "precip_3h"
        elif normalized in PRECIP_1H_NAMES:
            field = "precipitation"
        else:
            matched = [name for name, aliases in NORMALIZED_ALIASES.items() if normalized in aliases]
            if not matched:
                unknown.append(header)
                continue
            field = matched[0]
        pending.append((index, header, field))
    precip = [item for item in pending if item[2] == "precipitation"]
    if len(precip) > 1:
        day_items = [item for item in precip if normalize_header(item[1]) in PRECIP_DAY_NAMES]
        rest = [item for item in precip if item not in day_items]
        if len(rest) == 1 and day_items:
            day_keys = set(day_items)
            pending = [
                (index, header, "precip_day") if (index, header, field) in day_keys else (index, header, field)
                for index, header, field in pending
            ]
        else:
            raise WeatherImportError("'precipitation'에 매핑되는 헤더가 여러 개입니다.")
    mapping: dict[str, int] = {}
    for index, _header, field in pending:
        if field in mapping:
            raise WeatherImportError(f"'{field}'에 매핑되는 헤더가 여러 개입니다.")
        mapping[field] = index
    if "observation_date" not in mapping and "observation_datetime" not in mapping:
        raise WeatherImportError("관측일자 또는 관측일시 헤더를 인식하지 못했습니다.")
    if "station_id" not in mapping and "station_name" not in mapping:
        raise WeatherImportError("관측지점 번호 또는 관측지점명 헤더를 인식하지 못했습니다.")
    return mapping, unknown


def validate_file(path: Path) -> None:
    suffix = path.suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise WeatherImportError("XLSX, XLS, CSV 파일만 업로드할 수 있습니다.")
    if path.stat().st_size > MAX_FILE_SIZE:
        raise WeatherImportError(f"파일 크기는 {MAX_FILE_SIZE // (1024 * 1024)}MB 이하여야 합니다.")
    if suffix == ".xlsx":
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                    raise WeatherImportError("유효한 XLSX 파일이 아닙니다.")
                if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_SIZE:
                    raise WeatherImportError("압축 해제 크기가 제한을 초과합니다.")
                if archive.testzip():
                    raise WeatherImportError("손상된 XLSX 파일입니다.")
        except zipfile.BadZipFile as exc:
            raise WeatherImportError("유효한 XLSX 파일이 아닙니다.") from exc
    elif suffix == ".xls":
        if path.read_bytes()[:8] != bytes.fromhex("D0CF11E0A1B11AE1"):
            raise WeatherImportError("유효한 XLS 파일이 아닙니다.")
    else:
        sample = path.read_bytes()[:4096]
        if b"\x00" in sample:
            raise WeatherImportError("유효한 CSV 텍스트 파일이 아닙니다.")


def _decode_csv(path: Path) -> str:
    content = path.read_bytes()
    for encoding in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise WeatherImportError("CSV 인코딩은 UTF-8 또는 CP949여야 합니다.")


@contextmanager
def tabular_rows(path: Path) -> Iterator[Iterator[tuple[str, int, list[Any]]]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise WeatherImportError("XLSX 파일을 읽을 수 없습니다.") from exc
        try:
            def iterator() -> Iterator[tuple[str, int, list[Any]]]:
                for sheet in workbook.worksheets:
                    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                        yield sheet.title, row_number, list(values)
            yield iterator()
        finally:
            workbook.close()
        return
    if suffix == ".xls":
        try:
            import xlrd
            workbook = xlrd.open_workbook(path, on_demand=True)
        except Exception as exc:
            raise WeatherImportError("XLS 파일을 읽을 수 없습니다.") from exc
        try:
            def iterator() -> Iterator[tuple[str, int, list[Any]]]:
                for sheet in workbook.sheets():
                    for row_index in range(sheet.nrows):
                        values: list[Any] = []
                        for cell in sheet.row(row_index):
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(value, workbook.datemode)
                            values.append(value)
                        yield sheet.name, row_index + 1, values
            yield iterator()
        finally:
            workbook.release_resources()
        return
    text = _decode_csv(path)
    # weather.nuni.co.kr exports start with a "# source=KMA ..." metadata line; sniff on data lines only.
    sample = "\n".join(line for line in text[:16384].splitlines() if not line.lstrip().startswith("#"))
    try:
        dialect = csv.Sniffer().sniff(sample[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    yield (("CSV", number, row) for number, row in enumerate(csv.reader(text.splitlines(), dialect), start=1))


def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(float(value)) or float(value) != int(value):
            return None
        try:
            return date(1899, 12, 30) + timedelta(days=int(value))
        except OverflowError:
            return None
    text = str(value).strip()
    if " " in text:
        text = text.split(" ", 1)[0]
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y.%m.%d %H:%M", "%Y/%m/%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _parse_number(value: Any) -> tuple[float | None, str | None]:
    if value is None or str(value).strip().lower() in NULL_MARKERS:
        return None, None
    if isinstance(value, bool):
        return None, f'값 "{value}"을 숫자로 변환할 수 없습니다.'
    if isinstance(value, (int, float)):
        number = float(value)
        return (number, None) if math.isfinite(number) else (None, f'값 "{value}"을 숫자로 변환할 수 없습니다.')
    text = str(value).strip().replace(",", "")
    match = re.fullmatch(r"([-+]?(?:\d+(?:\.\d*)?|\.\d+))\s*(?:℃|°c|c|mm|cm|%|시간|hr|h)?", text, re.IGNORECASE)
    if not match:
        return None, f'값 "{value}"을 숫자로 변환할 수 없습니다.'
    return float(match.group(1)), None


def _station_value(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _name_station_id(name: str) -> str:
    normalized = re.sub(r"\s+", "", name).lower()
    if len(normalized) <= 64:
        return f"NAME:{normalized}"
    return f"NAME:{hashlib.sha256(normalized.encode('utf-8')).hexdigest()}"


def _fingerprint(rows: list[dict[str, Any]]) -> str:
    comparable = [{key: value for key, value in row.items() if key not in {"status", "existing", "raw_payload"}} for row in rows]
    raw = json.dumps(comparable, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _qc_code(value: Any) -> str | None:
    text = str(value or "").strip()
    return text if text in {"0", "1", "9"} else None


def _element_state(number: float | None, qc: str | None) -> tuple[float | None, str]:
    """QC 0/1/9 is documented for several ASOS elements. Blank is not treated as zero."""
    if qc == "1":
        return None, "error"
    if qc == "9":
        return None, "missing"
    if number is None:
        return None, "missing"
    return number, "value"


def _asos_headers(mapped_headers: dict[str, str]) -> bool:
    # Exact names only. 평균기온·일강수량·최심적설 are daily columns and must stay daily.
    names = {normalize_header(value) for value in mapped_headers.values()}
    return bool(names & {
        "기온", "1시간강수량", "1시간강수", "3시간강수량", "3시간강수",
        "3시간신적설", "현상번호", "일기현상", "기온qc플래그", "강수량qc플래그",
    })


def _explicit_clock(value: Any) -> datetime | None:
    """A real clock time. Midnight on an Excel date cell is not a clock."""
    if isinstance(value, datetime) and (value.hour or value.minute or value.second):
        return value.replace(tzinfo=None)
    text = str(value or "").strip()
    if ":" not in text:
        return None
    return _parse_datetime(text)


def _hub_layout(mapped_headers: dict[str, str]) -> bool:
    """weather.nuni 전체 레이아웃. ASOS 시간자료를 영문 열 이름으로 내려받은 파일이다."""
    names = {normalize_header(value) for value in mapped_headers.values()}
    return bool(names & {"snownew3h", "phenomenoncode", "rnqcflg", "taqcflg"})


def _seasonal_precip_interval(day: date | None) -> str:
    """data.go.kr ASOS rn. 11~3월은 3시간 간격, 4~10월은 1시간. 받은 파일로 확인했다."""
    if day is None:
        return "unconfirmed"
    return "3h" if day.month in {11, 12, 1, 2, 3} else "1h"


def _precip_interval(mapped_headers: dict[str, str], asos: bool) -> str:
    header = (mapped_headers.get("precipitation") or mapped_headers.get("precip_3h") or "").replace(" ", "")
    if "1시간강수" in header:
        return "1h"
    if "3시간강수" in header:
        return "3h"
    if asos:
        return ASOS_PRECIP_INTERVAL
    return "legacy_1h"


def _fields_absent_from_file(mapped_headers: dict[str, str]) -> set[str]:
    present = set(mapped_headers)
    if "precip_3h" in present:
        present.add("precipitation")
    skip: set[str] = set()
    for column, fields in _ABSENT_COLUMN_FIELDS.items():
        if column not in present:
            skip.update(fields)
    return skip


def parse_weather_file(path: Path, db: Session | None = None) -> dict[str, Any]:
    validate_file(path)
    mapping: dict[str, int] | None = None
    mapped_headers: dict[str, str] = {}
    unknown_headers: list[str] = []
    selected_sheet = ""
    header_row = 0
    header_cells: list[Any] = []
    rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: set[tuple[Any, str]] = set()
    duplicate_keys: set[tuple[Any, str]] = set()
    hourly_mode = False
    asos_mode = False
    hub_mode = False
    precip_interval = "legacy_1h"
    source_rows = 0
    source_note = ""
    first_seen: dict[tuple[Any, str], dict[str, Any]] = {}
    duplicate_rows = 0
    precipitation_blank_as_zero = 0
    outside_hour_rows = 0

    with tabular_rows(path) as source:
        for sheet_name, row_number, values in source:
            if mapping is None:
                first_cell = str(values[0] if values else "").strip()
                if first_cell.startswith("#") and not source_note:
                    source_note = ",".join(str(value) for value in values if value not in (None, "")).lstrip("# ").strip()[:300]
                if row_number > 30:
                    continue
                try:
                    candidate, unknown = map_headers(values)
                except WeatherImportError:
                    continue
                mapping, unknown_headers, selected_sheet, header_row = candidate, unknown, sheet_name, row_number
                header_cells = list(values)
                mapped_headers = {field: str(values[index] or "").strip() for field, index in candidate.items()}
                hourly_mode = "observation_datetime" in mapping
                asos_mode = _asos_headers(mapped_headers)
                hub_mode = _hub_layout(mapped_headers)
                precip_interval = "seasonal" if hub_mode else _precip_interval(mapped_headers, asos_mode)
                continue
            if sheet_name != selected_sheet or row_number <= header_row:
                continue
            if not any(value not in (None, "") for value in values):
                continue
            source_rows += 1
            if source_rows > MAX_ROWS:
                raise WeatherImportError(f"최대 {MAX_ROWS:,}행까지 업로드할 수 있습니다.")
            get = lambda field: values[mapping[field]] if field in mapping and mapping[field] < len(values) else None
            raw_time = get("observation_datetime") if "observation_datetime" in mapping else get("observation_date")
            clock = _explicit_clock(raw_time)
            if hourly_mode:
                observed_at = _parse_datetime(get("observation_datetime")) if "observation_datetime" in mapping else clock
                if observed_at:
                    observed_at = observed_at.replace(second=0, microsecond=0)
                    observed = observed_at.date()
                else:
                    observed = _parse_date(get("observation_date"))
            elif clock and asos_mode:
                hourly_mode = True
                observed_at = clock.replace(second=0, microsecond=0)
                observed = observed_at.date()
            else:
                observed_at = None
                observed = _parse_date(get("observation_date"))
            if hourly_mode and (asos_mode or hub_mode) and observed_at and observed_at.hour not in ASOS_KEEP_HOURS:
                outside_hour_rows += 1
                continue
            station_name = _station_value(get("station_name"))
            station_id = _station_value(get("station_id"))
            row_errors: list[str] = []
            warnings: list[str] = []
            if observed is None:
                row_errors.append("관측일시를 날짜와 시간으로 변환할 수 없습니다." if hourly_mode else "관측일자를 날짜로 변환할 수 없습니다.")
            if not station_id and not station_name:
                row_errors.append("관측지점 번호와 지점명이 모두 없습니다.")
            if not station_id and station_name:
                station_id = _name_station_id(station_name)
            fields = ("temperature", "precipitation", "humidity", "wind_speed", "snow_depth", "snow_new_3h") if hourly_mode else WEATHER_FIELDS
            weather: dict[str, float | None] = {}
            precipitation_filled = False
            for field in fields:
                if field not in mapping and field in {"snow_new_3h"}:
                    continue
                if field == "precipitation" and field not in mapping and "precip_3h" in mapping:
                    value, error = _parse_number(get("precip_3h"))
                else:
                    value, error = _parse_number(get(field))
                weather[field] = value
                if error:
                    warnings.append(f"{field}: {error} NULL로 처리합니다.")
            temp_qc = _qc_code(get("temp_qc")) if "temp_qc" in mapping else None
            precip_qc = _qc_code(get("precip_qc")) if "precip_qc" in mapping else None
            precip_mapped = "precipitation" in mapping or "precip_3h" in mapping
            if hourly_mode and (asos_mode or hub_mode):
                weather["temperature"], _temp_state = _element_state(weather.get("temperature"), temp_qc)
                if precip_mapped:
                    weather["precipitation"], precip_state = _element_state(weather.get("precipitation"), precip_qc)
                else:
                    precip_state = "absent"
                if "snow_depth" in mapping:
                    weather["snow_depth"], snow_state = _element_state(weather.get("snow_depth"), None)
                else:
                    snow_state = "absent"
                if "snow_new_3h" in mapping:
                    weather["snow_new_3h"], snow_new_state = _element_state(weather.get("snow_new_3h"), None)
                else:
                    snow_new_state = "absent"
            else:
                temp_state = "value" if weather.get("temperature") is not None else "missing"
                precip_state = "value" if weather.get("precipitation") is not None else "missing"
                snow_state = "value" if weather.get("snow_depth") is not None else ("absent" if "snow_depth" not in mapping else "missing")
                snow_new_state = "absent"
                if hourly_mode and "precipitation" in mapping and weather.get("precipitation") is None and not str(get("precipitation") or "").strip() \
                        and (weather.get("temperature") is not None or weather.get("humidity") is not None):
                    # 기존 weather.nuni 시간 CSV만. ASOS 공란은 0으로 바꾸지 않는다.
                    weather["precipitation"] = 0.0
                    precip_state = "legacy_blank"
                    precipitation_filled = True
                else:
                    precipitation_filled = False
                if hourly_mode and "snow_depth" in mapping and weather.get("snow_depth") is None and not str(get("snow_depth") or "").strip():
                    weather["snow_depth"] = 0.0
                    snow_state = "legacy_blank"
            if asos_mode or hub_mode:
                precipitation_filled = False
            phenomenon = str(get("phenomenon") or "").strip() if "phenomenon" in mapping else ""
            if phenomenon.endswith(".0") and phenomenon[:-2].isdigit():
                phenomenon = phenomenon[:-2]
            raw_payload = None
            if hourly_mode and (asos_mode or hub_mode) and header_cells:
                raw_payload = {}
                for index, header in enumerate(header_cells):
                    key = str(header or "").strip() or f"col{index}"
                    if index >= len(values):
                        continue
                    cell = values[index]
                    raw_payload[key] = cell.isoformat() if isinstance(cell, (datetime, date)) else cell
            identical_duplicate = False
            key_value = observed_at if hourly_mode else observed
            if key_value and station_id and not row_errors:
                key = (key_value, station_id)
                comparable = {
                    "station_name": station_name, **weather,
                    "source_kind": _station_value(get("source_kind")) if hourly_mode else "",
                    "precip_state": precip_state, "snow_state": snow_state, "snow_new_state": snow_new_state,
                }
                if key in seen:
                    if first_seen.get(key) == comparable:
                        identical_duplicate = True  # e.g. "01:00" and "01:00:00" rows with the same values
                        duplicate_rows += 1
                    else:
                        duplicate_keys.add(key)
                else:
                    seen.add(key)
                    first_seen[key] = comparable
            if identical_duplicate:
                continue
            precipitation_blank_as_zero += precipitation_filled
            row = {
                "source_row": row_number,
                "observation_date": observed.isoformat() if observed else "",
                "observation_datetime": observed_at.isoformat(timespec="minutes") if observed_at else "",
                "station_id": station_id,
                "station_name": station_name,
                **weather,
                "snow_new_3h": weather.get("snow_new_3h"),
                "precip_state": precip_state,
                "snow_state": snow_state,
                "snow_new_state": snow_new_state,
                "phenomenon": phenomenon,
                "precip_interval": (_seasonal_precip_interval(observed) if hub_mode else precip_interval) if hourly_mode and precip_mapped else None,
                "source_format": "hub_hourly" if hub_mode and hourly_mode else ("asos_hourly" if asos_mode and hourly_mode else ("hourly" if hourly_mode else "daily")),
                "raw_payload": raw_payload,
                "source_kind": _station_value(get("source_kind")) if hourly_mode else "",
                "warnings": warnings,
                "error": "; ".join(row_errors),
                "status": "오류" if row_errors else "신규",
            }
            rows.append(row)

    if mapping is None:
        raise WeatherImportError("파일의 처음 30행에서 필요한 헤더를 인식하지 못했습니다.")
    del first_seen
    for row in rows:
        key_text = row["observation_datetime"] if hourly_mode else row["observation_date"]
        if key_text:
            key_value = datetime.fromisoformat(key_text) if hourly_mode else date.fromisoformat(key_text)
            if (key_value, row["station_id"]) in duplicate_keys:
                label = "관측일시" if hourly_mode else "관측일자"
                row["error"] = "; ".join(filter(None, [row["error"], f"파일 안에서 {label}와 관측지점이 중복됩니다."]))
                row["status"] = "오류"
    errors = [{"row": row["source_row"], "message": row["error"]} for row in rows if row["error"]]

    valid_rows = [row for row in rows if not row["error"]]
    if db is not None and valid_rows:
        dates = [date.fromisoformat(row["observation_date"]) for row in valid_rows]
        station_ids = {row["station_id"] for row in valid_rows}
        if hourly_mode:
            datetimes = [datetime.fromisoformat(row["observation_datetime"]) for row in valid_rows]
            existing_rows = db.scalars(
                select(WeatherHourly).where(
                    WeatherHourly.observation_datetime.between(min(datetimes), max(datetimes)),
                    WeatherHourly.station_id.in_(station_ids),
                )
            ).all()
            existing_map = {(row.observation_datetime, row.station_id): row for row in existing_rows}
            skip = _fields_absent_from_file(mapped_headers)
            compare_fields = [field for field in HOURLY_COMPARE_FIELDS if field not in skip]
            for row in valid_rows:
                existing = existing_map.get((datetime.fromisoformat(row["observation_datetime"]), row["station_id"]))
                row["existing"] = existing.id if existing else None
                changed = existing and any(getattr(existing, field) != (None if row[field] == "" else row[field]) for field in compare_fields)
                row["status"] = "신규" if not existing else "수정" if changed else "변경 없음"
        else:
            existing_rows = db.scalars(
                select(WeatherHistory).where(
                    WeatherHistory.observation_date.between(min(dates), max(dates)),
                    WeatherHistory.station_id.in_(station_ids),
                )
            ).all()
            existing_map = {(row.observation_date, row.station_id): row for row in existing_rows}
            for row in valid_rows:
                existing = existing_map.get((date.fromisoformat(row["observation_date"]), row["station_id"]))
                row["existing"] = existing.id if existing else None
                row["status"] = "신규" if not existing else "수정" if any(getattr(existing, field) != row[field] for field in ("station_name", *WEATHER_FIELDS)) else "변경 없음"

    valid_dates = [date.fromisoformat(row["observation_date"]) for row in valid_rows]
    hours_per_day: dict[tuple[str, str], int] = {}
    if hourly_mode:
        for row in valid_rows:
            day_key = (row["observation_date"], row["station_id"])
            hours_per_day[day_key] = hours_per_day.get(day_key, 0) + 1
    stations = sorted({(row["station_id"], row["station_name"]) for row in valid_rows})
    summary = {
        "sheet_name": selected_sheet,
        "header_row": header_row,
        "data_granularity": "hourly" if hourly_mode else "daily",
        "mapped_headers": mapped_headers,
        "unknown_headers": unknown_headers,
        "total_rows": source_rows,
        "valid_rows": len(valid_rows),
        "inserted_rows": sum(row["status"] == "신규" for row in valid_rows),
        "updated_rows": sum(row["status"] == "수정" for row in valid_rows),
        "skipped_rows": sum(row["status"] == "변경 없음" for row in valid_rows),
        "error_rows": len(errors),
        "warning_fields": sum(len(row["warnings"]) for row in rows),
        "date_from": min(valid_dates).isoformat() if valid_dates else None,
        "date_to": max(valid_dates).isoformat() if valid_dates else None,
        "stations": [{"station_id": station_id, "station_name": name} for station_id, name in stations],
        "rows_fingerprint": _fingerprint(valid_rows),
        "source_note": source_note,
        "duplicate_rows": duplicate_rows,
        "precipitation_blank_as_zero": precipitation_blank_as_zero,
        "outside_hour_rows": outside_hour_rows,
        "file_format": "hub_hourly" if hub_mode and hourly_mode else ("asos_hourly" if asos_mode and hourly_mode else ("hourly" if hourly_mode else "daily")),
        "precip_interval": precip_interval if hourly_mode else "daily",
        "missing_temperature": sum(row.get("temperature") is None for row in valid_rows),
        "missing_precipitation": sum(row.get("precipitation") is None for row in valid_rows),
        "missing_snow": sum(row.get("snow_depth") is None for row in valid_rows) if "snow_depth" in mapped_headers else None,
        "missing_snow_new": sum(row.get("snow_new_3h") is None for row in valid_rows) if "snow_new_3h" in mapped_headers else None,
        "daily_aggregate_days": sum(count >= DAILY_MIN_HOURLY_SAMPLES for count in hours_per_day.values()),
        "partial_days": sum(count < DAILY_MIN_HOURLY_SAMPLES for count in hours_per_day.values()),
        "daily_min_hourly_samples": DAILY_MIN_HOURLY_SAMPLES if hourly_mode else None,
    }
    return {"summary": summary, "rows": rows, "errors": errors}


def _refresh_meal_period_weather(db: Session, rows: list[dict[str, Any]], batch_id: int, now: datetime) -> None:
    affected = {(date.fromisoformat(row["observation_date"]), row["station_id"]) for row in rows}
    if not affected:
        return
    dates = [item[0] for item in affected]
    station_ids = {item[1] for item in affected}
    start_at = datetime.combine(min(dates), datetime.min.time())
    end_at = datetime.combine(max(dates) + timedelta(days=1), datetime.min.time())
    hourly_rows = db.scalars(
        select(WeatherHourly).where(
            WeatherHourly.observation_datetime >= start_at,
            WeatherHourly.observation_datetime < end_at,
            WeatherHourly.station_id.in_(station_ids),
        )
    ).all()
    by_day: dict[tuple[date, str], dict[int, WeatherHourly]] = {}
    for row in hourly_rows:
        key = (row.observation_datetime.date(), row.station_id)
        if key not in affected:
            continue
        by_day.setdefault(key, {})[row.observation_datetime.hour] = row
    values = []
    for (observation_date, station_id), by_hour in by_day.items():
        for meal_type, meal_hours in MEAL_WEATHER_HOURS.items():
            samples = [by_hour[hour] for hour in meal_hours if hour in by_hour]
            if not samples:
                continue

            def numbers(field: str, chosen: list[WeatherHourly] = samples) -> list[float]:
                return [float(value) for sample in chosen if (value := getattr(sample, field)) is not None]

            temperatures = numbers("temperature")
            precipitation = numbers("precipitation")
            humidity = numbers("humidity")
            wind_speed = numbers("wind_speed")
            snow = numbers("snow_depth")
            asos = any(getattr(sample, "source_format", None) in {"asos_hourly", "hub_hourly"} for sample in samples)
            if asos:
                end_hour = max(meal_hours)
                temp_rows = [by_hour.get(hour) for hour in meal_hours]
                temp_values = [float(row.temperature) for row in temp_rows if row is not None and row.temperature is not None]
                temp_complete = len(temp_values) == len(meal_hours)
                intervals = {sample.precip_interval for sample in samples if sample.precip_interval}
                interval = next(iter(intervals)) if len(intervals) == 1 else (ASOS_PRECIP_INTERVAL if not intervals else "unconfirmed")
                if interval == "1h":
                    end = by_hour.get(end_hour)
                    if end is not None and end.precip_state == "value" and end.precipitation is not None:
                        precip_amount, precip_state = float(end.precipitation), "ok"
                    else:
                        precip_amount, precip_state = None, "unavailable"
                    pre_rows = [by_hour.get(hour) for hour in PRE_MEAL_HOURS[meal_type]]
                    if all(row is not None and row.precip_state == "value" and row.precipitation is not None for row in pre_rows):
                        pre_amount, pre_state = sum(float(row.precipitation) for row in pre_rows), "ok"
                    else:
                        pre_amount, pre_state = None, "unavailable"
                    method = "1h_end"
                else:
                    precip_amount, precip_state, pre_amount, pre_state, method = None, "unavailable", None, "unavailable", interval
                snow_rows = [by_hour.get(hour) for hour in meal_hours]
                snow_known = [row for row in snow_rows if row is not None and row.snow_state == "value" and row.snow_depth is not None]
                if len(snow_known) == len(meal_hours):
                    snow_amount, snow_state = max(float(row.snow_depth) for row in snow_known), "ok"
                elif snow_known:
                    snow_amount, snow_state = max(float(row.snow_depth) for row in snow_known), "partial"
                else:
                    snow_amount, snow_state = None, "missing"
                fresh = by_hour.get(end_hour)
                if fresh is not None and fresh.snow_new_state == "value" and fresh.snow_new_3h is not None:
                    fresh_amount, fresh_state = float(fresh.snow_new_3h), "value"
                elif fresh is not None and fresh.snow_new_state == "missing":
                    fresh_amount, fresh_state = None, "missing"
                else:
                    fresh_amount, fresh_state = None, "absent"
                values.append({
                    "observation_date": observation_date,
                    "station_id": station_id,
                    "station_name": next((sample.station_name for sample in samples if sample.station_name), None),
                    "meal_type": meal_type,
                    "avg_temp": sum(temp_values) / len(temp_values) if temp_values else None,
                    "min_temp": min(temp_values) if temp_values else None,
                    "max_temp": max(temp_values) if temp_values else None,
                    "precipitation": precip_amount,
                    "avg_humidity": sum(humidity) / len(humidity) if humidity else None,
                    "avg_wind_speed": sum(wind_speed) / len(wind_speed) if wind_speed else None,
                    "snow_depth": snow_amount,
                    "snow_state": snow_state,
                    "snow_new_3h": fresh_amount,
                    "snow_new_state": fresh_state,
                    "precip_state": precip_state,
                    "precip_method": method,
                    "pre_precip": pre_amount,
                    "pre_precip_state": pre_state,
                    "temp_complete": temp_complete,
                    "interpretation": INTERPRETATION_VERSION,
                    "sample_count": len(samples),
                    "upload_batch_id": batch_id,
                    "created_at": now,
                    "updated_at": now,
                })
                continue
            legacy_precip = "legacy_blank" if any(sample.precip_state == "legacy_blank" for sample in samples) else ("legacy" if precipitation else None)
            legacy_snow = "legacy_blank" if any(sample.snow_state == "legacy_blank" for sample in samples) else ("legacy" if snow else None)
            values.append({
                "observation_date": observation_date,
                "station_id": station_id,
                "station_name": next((sample.station_name for sample in samples if sample.station_name), None),
                "meal_type": meal_type,
                "avg_temp": sum(temperatures) / len(temperatures) if temperatures else None,
                "min_temp": min(temperatures) if temperatures else None,
                "max_temp": max(temperatures) if temperatures else None,
                "precipitation": sum(precipitation) if precipitation else None,
                "avg_humidity": sum(humidity) / len(humidity) if humidity else None,
                "avg_wind_speed": sum(wind_speed) / len(wind_speed) if wind_speed else None,
                "snow_depth": max(snow) if snow else None,
                "snow_state": legacy_snow,
                "snow_new_3h": None,
                "snow_new_state": None,
                "precip_state": legacy_precip,
                "precip_method": "legacy_sum",
                "pre_precip": None,
                "pre_precip_state": None,
                "temp_complete": len(temperatures) == len(meal_hours),
                "interpretation": "legacy_hourly",
                "sample_count": len(samples),
                "upload_batch_id": batch_id,
                "created_at": now,
                "updated_at": now,
            })
    if not values:
        return
    insert = pg_insert(MealPeriodWeather) if db.bind and db.bind.dialect.name == "postgresql" else sqlite_insert(MealPeriodWeather)
    statement = insert.values(values)
    statement = statement.on_conflict_do_update(
        index_elements=[MealPeriodWeather.observation_date, MealPeriodWeather.station_id, MealPeriodWeather.meal_type],
        set_={
            "station_name": statement.excluded.station_name,
            "avg_temp": statement.excluded.avg_temp,
            "min_temp": statement.excluded.min_temp,
            "max_temp": statement.excluded.max_temp,
            "precipitation": statement.excluded.precipitation,
            "avg_humidity": statement.excluded.avg_humidity,
            "avg_wind_speed": statement.excluded.avg_wind_speed,
            "snow_depth": statement.excluded.snow_depth,
            "snow_state": statement.excluded.snow_state,
            "snow_new_3h": statement.excluded.snow_new_3h,
            "snow_new_state": statement.excluded.snow_new_state,
            "precip_state": statement.excluded.precip_state,
            "precip_method": statement.excluded.precip_method,
            "pre_precip": statement.excluded.pre_precip,
            "pre_precip_state": statement.excluded.pre_precip_state,
            "temp_complete": statement.excluded.temp_complete,
            "interpretation": statement.excluded.interpretation,
            "sample_count": statement.excluded.sample_count,
            "upload_batch_id": statement.excluded.upload_batch_id,
            "updated_at": statement.excluded.updated_at,
        },
    )
    db.execute(statement)


def _refresh_daily_weather_from_hourly(db: Session, rows: list[dict[str, Any]], batch_id: int, now: datetime) -> int:
    """Upsert weather_history daily rows (source KMA_HOURLY) from all stored hourly rows of the affected days.

    Days with fewer than DAILY_MIN_HOURLY_SAMPLES hours are skipped, and a daily row that came from an
    official daily file (source KMA_FILE, which also has snow/sunshine) is never overwritten.
    """
    affected = {(date.fromisoformat(row["observation_date"]), row["station_id"]) for row in rows}
    if not affected:
        return 0
    dates = [item[0] for item in affected]
    station_ids = {item[1] for item in affected}
    hourly_rows = db.scalars(
        select(WeatherHourly).where(
            WeatherHourly.observation_datetime >= datetime.combine(min(dates), datetime.min.time()),
            WeatherHourly.observation_datetime < datetime.combine(max(dates) + timedelta(days=1), datetime.min.time()),
            WeatherHourly.station_id.in_(station_ids),
        )
    ).all()
    official = {
        (row.observation_date, row.station_id)
        for row in db.execute(
            select(WeatherHistory.observation_date, WeatherHistory.station_id).where(
                WeatherHistory.observation_date.between(min(dates), max(dates)),
                WeatherHistory.station_id.in_(station_ids),
                WeatherHistory.source == "KMA_FILE",
            )
        )
    }
    grouped: dict[tuple[date, str], list[WeatherHourly]] = {}
    for row in hourly_rows:
        key = (row.observation_datetime.date(), row.station_id)
        if key in affected and key not in official:
            grouped.setdefault(key, []).append(row)
    values = []
    for (observation_date, station_id), samples in grouped.items():
        if any(getattr(sample, "source_format", None) in {"asos_hourly", "hub_hourly"} for sample in samples):
            continue
        if len(samples) < DAILY_MIN_HOURLY_SAMPLES:
            continue

        def numbers(field: str) -> list[float]:
            return [float(value) for sample in samples if (value := getattr(sample, field)) is not None]

        temperatures, precipitation, humidity = numbers("temperature"), numbers("precipitation"), numbers("humidity")
        values.append({
            "observation_date": observation_date,
            "station_id": station_id,
            "station_name": next((sample.station_name for sample in samples if sample.station_name), None),
            "avg_temp": round(sum(temperatures) / len(temperatures), 1) if temperatures else None,
            "min_temp": min(temperatures) if temperatures else None,
            "max_temp": max(temperatures) if temperatures else None,
            "precipitation": round(sum(precipitation), 1) if precipitation else None,
            "avg_humidity": round(sum(humidity) / len(humidity), 1) if humidity else None,
            "snow_depth": None,
            "sunshine_hours": None,
            "source": "KMA_HOURLY",
            "upload_batch_id": batch_id,
            "created_at": now,
            "updated_at": now,
        })
    if not values:
        return 0
    insert = pg_insert(WeatherHistory) if db.bind and db.bind.dialect.name == "postgresql" else sqlite_insert(WeatherHistory)
    for start in range(0, len(values), 1000):
        statement = insert.values(values[start:start + 1000])
        statement = statement.on_conflict_do_update(
            index_elements=[WeatherHistory.observation_date, WeatherHistory.station_id],
            set_={
                "station_name": statement.excluded.station_name,
                **{field: getattr(statement.excluded, field) for field in WEATHER_FIELDS},
                "source": statement.excluded.source,
                "upload_batch_id": statement.excluded.upload_batch_id,
                "updated_at": statement.excluded.updated_at,
            },
            where=WeatherHistory.source != "KMA_FILE",
        )
        db.execute(statement)
    return len(values)


def apply_weather_file(path: Path, db: Session, filename: str, checksum: str, user_id: int, expected_fingerprint: str) -> tuple[WeatherUploadHistory, dict[str, int]]:
    parsed = parse_weather_file(path, db)
    if parsed["summary"]["rows_fingerprint"] != expected_fingerprint:
        raise WeatherImportError("Preview 이후 파일 또는 기존 날씨자료가 변경되었습니다. 다시 분석해 주세요.")
    summary = parsed["summary"]
    batch = WeatherUploadHistory(
        original_filename=filename,
        checksum_sha256=checksum,
        date_from=date.fromisoformat(summary["date_from"]) if summary["date_from"] else None,
        date_to=date.fromisoformat(summary["date_to"]) if summary["date_to"] else None,
        stations=summary["stations"],
        total_rows=summary["total_rows"],
        valid_rows=summary["valid_rows"],
        inserted_rows=summary["inserted_rows"],
        updated_rows=summary["updated_rows"],
        skipped_rows=summary["skipped_rows"],
        error_rows=summary["error_rows"],
        status="PROCESSING",
        errors=parsed["errors"][:1000],
        uploaded_by=user_id,
    )
    db.add(batch)
    db.flush()
    now = datetime.now(timezone.utc)
    changed = [row for row in parsed["rows"] if row["status"] in {"신규", "수정"}]
    daily_refreshed = 0
    if changed:
        if summary["data_granularity"] == "hourly":
            values = [{
                "observation_datetime": datetime.fromisoformat(row["observation_datetime"]),
                "station_id": row["station_id"],
                "station_name": row["station_name"] or None,
                "temperature": row["temperature"],
                "precipitation": row["precipitation"],
                "humidity": row["humidity"],
                "wind_speed": row["wind_speed"],
                "snow_depth": row.get("snow_depth"),
                "snow_state": row.get("snow_state"),
                "snow_new_3h": row.get("snow_new_3h"),
                "snow_new_state": row.get("snow_new_state"),
                "precip_state": row.get("precip_state"),
                "precip_interval": row.get("precip_interval"),
                "phenomenon": row.get("phenomenon") or None,
                "source_format": row.get("source_format"),
                "raw_payload": row.get("raw_payload"),
                "source_kind": row["source_kind"] or None,
                "upload_batch_id": batch.id,
                "created_at": now,
                "updated_at": now,
            } for row in changed]
            mapped = summary.get("mapped_headers") or {}
            skip = _fields_absent_from_file(mapped)
            if skip:
                moments = [item["observation_datetime"] for item in values]
                existing_rows = db.scalars(
                    select(WeatherHourly).where(
                        WeatherHourly.observation_datetime.between(min(moments), max(moments)),
                        WeatherHourly.station_id.in_({item["station_id"] for item in values}),
                    )
                ).all()
                existing_map = {(row.observation_datetime, row.station_id): row for row in existing_rows}
                for item in values:
                    current = existing_map.get((item["observation_datetime"], item["station_id"]))
                    if not current:
                        continue
                    for name in skip:
                        if hasattr(current, name):
                            item[name] = getattr(current, name)
            insert = pg_insert(WeatherHourly) if db.bind and db.bind.dialect.name == "postgresql" else sqlite_insert(WeatherHourly)
            for start in range(0, len(values), 1000):
                statement = insert.values(values[start:start + 1000])
                statement = statement.on_conflict_do_update(
                    index_elements=[WeatherHourly.observation_datetime, WeatherHourly.station_id],
                    set_={
                        "station_name": statement.excluded.station_name,
                        "temperature": statement.excluded.temperature,
                        "precipitation": statement.excluded.precipitation,
                        "humidity": statement.excluded.humidity,
                        "wind_speed": statement.excluded.wind_speed,
                        "snow_depth": statement.excluded.snow_depth,
                        "snow_state": statement.excluded.snow_state,
                        "snow_new_3h": statement.excluded.snow_new_3h,
                        "snow_new_state": statement.excluded.snow_new_state,
                        "precip_state": statement.excluded.precip_state,
                        "precip_interval": statement.excluded.precip_interval,
                        "phenomenon": statement.excluded.phenomenon,
                        "source_format": statement.excluded.source_format,
                        "raw_payload": statement.excluded.raw_payload,
                        "source_kind": statement.excluded.source_kind,
                        "upload_batch_id": statement.excluded.upload_batch_id,
                        "updated_at": statement.excluded.updated_at,
                    },
                )
                db.execute(statement)
            _refresh_meal_period_weather(db, changed, batch.id, now)
            daily_refreshed = _refresh_daily_weather_from_hourly(db, changed, batch.id, now)
        else:
            values = [{
                "observation_date": date.fromisoformat(row["observation_date"]),
                "station_id": row["station_id"],
                "station_name": row["station_name"] or None,
                **{field: row[field] for field in WEATHER_FIELDS},
                "source": "KMA_FILE",
                "upload_batch_id": batch.id,
                "created_at": now,
                "updated_at": now,
            } for row in changed]
            insert = pg_insert(WeatherHistory) if db.bind and db.bind.dialect.name == "postgresql" else sqlite_insert(WeatherHistory)
            for start in range(0, len(values), 1000):
                statement = insert.values(values[start:start + 1000])
                statement = statement.on_conflict_do_update(
                    index_elements=[WeatherHistory.observation_date, WeatherHistory.station_id],
                    set_={
                        "station_name": statement.excluded.station_name,
                        **{field: getattr(statement.excluded, field) for field in WEATHER_FIELDS},
                        "source": statement.excluded.source,
                        "upload_batch_id": statement.excluded.upload_batch_id,
                        "updated_at": statement.excluded.updated_at,
                    },
                )
                db.execute(statement)
    batch.status = "COMPLETED_WITH_ERRORS" if summary["error_rows"] else "COMPLETED"
    batch.completed_at = now
    result = {
        "total_rows": summary["total_rows"],
        "inserted_rows": summary["inserted_rows"],
        "updated_rows": summary["updated_rows"],
        "skipped_rows": summary["skipped_rows"],
        "error_rows": summary["error_rows"],
        "duplicate_rows": summary.get("duplicate_rows", 0),
        "daily_aggregated_days": daily_refreshed,
    }
    return batch, result
