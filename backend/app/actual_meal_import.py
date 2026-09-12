from __future__ import annotations

import hashlib
import json
import math
import re
import zipfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import MealActual, MealService, MealTypeSetting

EXPECTED_FILENAME = "실제식수정보_통합(1).xlsx"
EXPECTED_SHEET = "실제식수정보"
REQUIRED_HEADERS = ("일자", "중식", "석식", "특이사항")
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_ZIP_FILES = 1000
PREVIEW_TTL = timedelta(minutes=30)


class ActualMealUploadError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_xlsx_container(path: Path) -> None:
    size = path.stat().st_size
    if size > MAX_FILE_SIZE:
        raise ActualMealUploadError(f"파일 크기는 {MAX_FILE_SIZE // (1024 * 1024)}MB 이하여야 합니다.")
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ActualMealUploadError("유효한 XLSX 파일이 아닙니다.")
            if len(names) > MAX_ZIP_FILES:
                raise ActualMealUploadError("압축 파일 내부 항목이 너무 많습니다.")
            total = sum(info.file_size for info in archive.infolist())
            if total > MAX_UNCOMPRESSED_SIZE:
                raise ActualMealUploadError("압축 해제 파일 크기가 제한을 초과합니다.")
            bad = archive.testzip()
            if bad:
                raise ActualMealUploadError("손상된 XLSX 파일입니다.")
    except zipfile.BadZipFile as exc:
        raise ActualMealUploadError("유효한 XLSX 파일이 아닙니다.") from exc


def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(float(value)) or float(value) != int(value):
            return None
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_count(value: Any) -> tuple[int | None, str | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, bool):
        return None, "중식·석식은 정수만 입력할 수 있습니다."
    if isinstance(value, int):
        count = value
    elif isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return None, "소수는 입력할 수 없습니다."
        count = int(value)
    else:
        text = str(value).strip()
        if not re.fullmatch(r"[+-]?\d+", text):
            return None, "숫자로 변환할 수 없는 값입니다."
        try:
            count = int(Decimal(text))
        except (InvalidOperation, ValueError):
            return None, "숫자로 변환할 수 없는 값입니다."
    if count < 0:
        return None, "음수는 입력할 수 없습니다."
    return count, None


def _row_fingerprint(rows: list[dict[str, Any]]) -> str:
    payload = json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _meal_type_map(db: Session) -> dict[str, MealTypeSetting]:
    settings = db.scalars(select(MealTypeSetting).where(MealTypeSetting.active.is_(True))).all()
    result: dict[str, MealTypeSetting] = {}
    for item in settings:
        result[item.name] = item
        result[item.code] = item
    return result


def preview_actual_meals(path: Path, db: Session) -> dict[str, Any]:
    validate_xlsx_container(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ActualMealUploadError("XLSX 파일을 읽을 수 없습니다.") from exc

    try:
        if workbook.sheetnames != [EXPECTED_SHEET]:
            raise ActualMealUploadError(f"시트는 '{EXPECTED_SHEET}' 1개만 있어야 합니다.")
        sheet = workbook[EXPECTED_SHEET]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_values = next(iterator)
        except StopIteration as exc:
            raise ActualMealUploadError("헤더 행이 없습니다.") from exc
        headers = [str(value).strip() if value is not None else "" for value in header_values]
        if len(headers) != len(set(headers)) or set(headers) != set(REQUIRED_HEADERS):
            raise ActualMealUploadError("필수 열은 일자, 중식, 석식, 특이사항이어야 합니다.")
        positions = {name: headers.index(name) for name in REQUIRED_HEADERS}

        type_map = _meal_type_map(db)
        rows: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        source_row_count = 0
        dates: list[date] = []
        duplicate_keys: set[tuple[date, str]] = set()
        seen_keys: set[tuple[date, str]] = set()
        lunch_sum = dinner_sum = 0
        lunch_count = dinner_count = 0
        excluded_count = 0

        for excel_row, values in enumerate(iterator, start=2):
            cells = list(values)
            if not any(value not in (None, "") for value in cells):
                continue
            source_row_count += 1
            raw_date = cells[positions["일자"]] if positions["일자"] < len(cells) else None
            parsed_date = _parse_date(raw_date)
            note_value = cells[positions["특이사항"]] if positions["특이사항"] < len(cells) else ""
            note = str(note_value).strip() if note_value not in (None, "") else ""
            row_errors: list[str] = []
            if parsed_date is None:
                row_errors.append("존재하지 않는 날짜입니다.")
            else:
                dates.append(parsed_date)
            parsed_counts: dict[str, int | None] = {}
            count_errors: dict[str, str | None] = {}
            for column, type_name in (("중식", "중식"), ("석식", "석식")):
                value = cells[positions[column]] if positions[column] < len(cells) else None
                count, error = _parse_count(value)
                if value in (None, ""):
                    excluded_count += 1
                parsed_counts[type_name] = count
                count_errors[type_name] = error
                if error:
                    row_errors.append(f"{column}: {error}")
                if parsed_date is not None and value not in (None, ""):
                    key = (parsed_date, type_name)
                    if key in seen_keys:
                        duplicate_keys.add(key)
                    seen_keys.add(key)
                if count is not None:
                    if column == "중식":
                        lunch_count += 1
                        lunch_sum += count
                    else:
                        dinner_count += 1
                        dinner_sum += count
            for type_name, count in parsed_counts.items():
                if count is None and not count_errors[type_name]:
                    continue
                setting = type_map.get(type_name)
                if not setting:
                    row_errors.append(f"활성 배식유형 '{type_name}'을 찾을 수 없습니다.")
                    continue
                key = (parsed_date, type_name) if parsed_date else None
                if key in duplicate_keys:
                    row_errors.append(f"날짜와 배식유형이 중복됩니다: {parsed_date} {type_name}")
                row = {
                    "excel_row": excel_row,
                    "date": parsed_date.isoformat() if parsed_date else "",
                    "meal_type": setting.code,
                    "meal_type_name": setting.name,
                    "upload_count": count,
                    "note": note,
                    "error": "; ".join(row_errors),
                }
                rows.append(row)

        # Duplicate detection is completed after reading all rows so both sides are marked.
        for row in rows:
            key = (date.fromisoformat(row["date"]), row["meal_type_name"]) if row["date"] else None
            if key in duplicate_keys:
                message = f"날짜와 배식유형이 중복됩니다: {row['date']} {row['meal_type_name']}"
                row["error"] = "; ".join(filter(None, [row["error"], message]))

        dates_for_query = {date.fromisoformat(row["date"]) for row in rows if row["date"]}
        services = db.scalars(select(MealService).where(MealService.service_date.in_(dates_for_query))).all() if dates_for_query else []
        service_map = {(item.service_date, item.meal_type): item for item in services}
        for row in rows:
            if row["error"]:
                row.update({"existing_count": None, "status": "오류", "service_id": None})
                continue
            service = service_map.get((date.fromisoformat(row["date"]), row["meal_type"]))
            if not service:
                setting = type_map.get(row["meal_type_name"])
                if not setting:
                    row.update({"existing_count": None, "status": "오류", "service_id": None})
                    row["error"] = f"활성 배식유형 '{row['meal_type_name']}'을 찾을 수 없습니다."
                    continue
                service = MealService(
                    service_date=date.fromisoformat(row["date"]),
                    meal_type=setting.code,
                    planned_count=setting.default_planned_count,
                    service_time=setting.default_service_time,
                )
                db.add(service)
                db.flush()
                service_map[(service.service_date, service.meal_type)] = service
                row["service_created"] = True
            else:
                row["service_created"] = False
            existing = service.actual
            row["service_id"] = service.id
            row["existing_count"] = existing.actual_count if existing else None
            row["existing_note"] = existing.note if existing else ""
            if not existing or existing.actual_count is None:
                row["status"] = "신규"
            elif existing.actual_count != row["upload_count"] or (existing.note or "") != row["note"]:
                row["status"] = "수정"
            else:
                row["status"] = "변경 없음"

        error_count = sum(1 for row in rows if row["status"] == "오류")
        summary = {
            "sheet_count": 1,
            "sheet_name": EXPECTED_SHEET,
            "source_row_count": source_row_count,
            "start_date": min(dates).isoformat() if dates else None,
            "end_date": max(dates).isoformat() if dates else None,
            "lunch_input_count": lunch_count,
            "lunch_sum": lunch_sum,
            "dinner_input_count": dinner_count,
            "dinner_sum": dinner_sum,
            "candidate_count": len(rows),
            "new_count": sum(1 for row in rows if row["status"] == "신규"),
            "update_count": sum(1 for row in rows if row["status"] == "수정"),
            "unchanged_count": sum(1 for row in rows if row["status"] == "변경 없음"),
            "service_created_count": sum(1 for row in rows if row.get("service_created")),
            "excluded_count": excluded_count,
            "error_count": error_count,
            "rows_fingerprint": _row_fingerprint(rows),
        }
        return {"summary": summary, "rows": rows, "errors": [{"excel_row": row["excel_row"], "message": row["error"]} for row in rows if row["error"]]}
    finally:
        workbook.close()


def apply_actual_meals(path: Path, db: Session, user_id: int, expected_fingerprint: str) -> dict[str, int]:
    parsed = preview_actual_meals(path, db)
    if parsed["summary"]["rows_fingerprint"] != expected_fingerprint:
        raise ActualMealUploadError("Preview 이후 파일 또는 DB 연결 정보가 변경되었습니다. 다시 검증해 주세요.")
    if parsed["errors"]:
        raise ActualMealUploadError("오류가 있는 Preview는 반영할 수 없습니다.")
    now = datetime.now(timezone.utc)
    result = {"candidate_count": len(parsed["rows"]), "new_count": 0, "update_count": 0, "unchanged_count": 0, "service_created_count": 0, "excluded_count": parsed["summary"]["excluded_count"], "failed_count": 0}
    for row in parsed["rows"]:
        service = db.get(MealService, row["service_id"])
        if not service:
            setting = db.scalar(select(MealTypeSetting).where(MealTypeSetting.code == row["meal_type"], MealTypeSetting.active.is_(True)))
            if not setting:
                raise ActualMealUploadError(f"활성 배식유형 '{row['meal_type_name']}'을 찾을 수 없습니다: {row['date']}")
            service = MealService(
                service_date=date.fromisoformat(row["date"]),
                meal_type=setting.code,
                planned_count=setting.default_planned_count,
                service_time=setting.default_service_time,
            )
            db.add(service)
            db.flush()
            result["service_created_count"] += 1
        actual = service.actual
        if actual and actual.actual_count == row["upload_count"] and (actual.note or "") == row["note"]:
            result["unchanged_count"] += 1
            continue
        if actual is None:
            actual = MealActual(meal_service_id=service.id)
            result["new_count"] += 1
        else:
            result["update_count"] += 1
        actual.actual_count = row["upload_count"]
        actual.note = row["note"] or None
        actual.recorded_at = now
        db.add(actual)
    return result
