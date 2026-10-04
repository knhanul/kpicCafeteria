from __future__ import annotations

import io
from datetime import date
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from .. import meal_count_analysis as analysis
from ..db import get_db
from ..deps import current_user
from ..models import User

router = APIRouter(prefix="/api/analysis", tags=["analysis"])


def _run(fn, *args, **kwargs):
    try:
        return fn(*args, **kwargs)
    except analysis.AnalysisError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/daily")
def daily(start: date, end: date, meal_type: str = "LUNCH", db: Session = Depends(get_db), user: User = Depends(current_user)):
    return _run(analysis.daily, db, meal_type, start, end)


@router.get("/menu-groups/search")
def menu_groups_search(q: str = "", db: Session = Depends(get_db), user: User = Depends(current_user)):
    return {"items": analysis.search_menu_groups(db, q)}


@router.get("/menu-group")
def menu_group(
    start: date, end: date, name: str, meal_type: str = "LUNCH", weather: str = "all",
    db: Session = Depends(get_db), user: User = Depends(current_user),
):
    return _run(analysis.menu_group, db, meal_type, start, end, name, weather)


@router.get("/menus/search")
def menus_search(q: str = "", mode: str = "group", db: Session = Depends(get_db), user: User = Depends(current_user)):
    return {"items": _run(analysis.search_menus, db, mode, q)}


@router.get("/menus")
def menus_compare(
    start: date, end: date, mode: str = "group", names: list[str] = Query(default=[]), meal_type: str = "LUNCH",
    weather: str = "all", db: Session = Depends(get_db), user: User = Depends(current_user),
):
    return _run(analysis.menu_compare, db, meal_type, start, end, mode, names, weather)


@router.get("/popular-menus")
def popular_menus(
    start: date, end: date, meal_type: str = "LUNCH", basis: str = "name", order: str = "top", limit: int = 10,
    min_days: int | None = None, main_only: bool = False,
    db: Session = Depends(get_db), user: User = Depends(current_user),
):
    return _run(analysis.popular_menus, db, meal_type, start, end, basis, order, limit, min_days, main_only)


@router.get("/ingredients/search")
def ingredients_search(q: str = "", db: Session = Depends(get_db), user: User = Depends(current_user)):
    return {"items": analysis.search_ingredients(db, q)}


@router.get("/ingredient")
def ingredient(
    start: date, end: date, ingredient_id: int, meal_type: str = "LUNCH", weather: str = "all",
    db: Session = Depends(get_db), user: User = Depends(current_user),
):
    return _run(analysis.ingredient, db, meal_type, start, end, ingredient_id, weather)


@router.get("/weather")
def weather(start: date, end: date, meal_type: str = "LUNCH", db: Session = Depends(get_db), user: User = Depends(current_user)):
    return _run(analysis.weather_bands, db, meal_type, start, end)


@router.get("/date-detail")
def date_detail(
    service_date: date = Query(alias="date"), meal_type: str = "LUNCH", ingredient_id: int | None = None,
    db: Session = Depends(get_db), user: User = Depends(current_user),
):
    return _run(analysis.date_detail, db, meal_type, service_date, ingredient_id)


def _point_rows(points: list[dict[str, Any]], meal_name: str) -> list[list[Any]]:
    return [
        [
            p["date"], p["weekday"], meal_name, p["actual"], p["usual"], p["diff"],
            (p["weather"] or {}).get("text") or "",
        ]
        for p in points
    ]


POINT_HEADERS = ["날짜", "요일", "배식", "실제(명)", "평소(명)", "차이(명)", "배식시간 날씨"]


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for index, width in enumerate([12, 6, 8, 10, 10, 10, 30][: len(headers)], start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    return ws


@router.get("/export.xlsx")
def export_xlsx(
    tab: str, start: date, end: date, meal_type: str = "LUNCH", name: str = "", ingredient_id: int | None = None,
    weather: str = "all", mode: str = "group", names: list[str] = Query(default=[]),
    basis: str = "name", order: str = "top", limit: int = 10, min_days: int | None = None, main_only: bool = False,
    db: Session = Depends(get_db), user: User = Depends(current_user),
):
    wb = Workbook()
    meal_name = analysis.MEAL_TYPES.get(meal_type, meal_type)
    if tab == "daily":
        data = _run(analysis.daily, db, meal_type, start, end)
        label = "날짜별식수"
        _sheet(wb, label, POINT_HEADERS, _point_rows(data["points"], meal_name), first=True)
    elif tab == "menu":
        data = _run(analysis.menu_group, db, meal_type, start, end, name, weather)
        label = f"메뉴별식수_{data['menu_group']}"
        _sheet(wb, "메뉴별식수", POINT_HEADERS, _point_rows(data["points"], meal_name), first=True)
    elif tab == "menus":
        data = _run(analysis.menu_compare, db, meal_type, start, end, mode, names, weather)
        label = "메뉴비교_" + "_".join(item["name"] for item in data["items"])[:60]
        rows = [[item["name"], *row] for item in data["items"] for row in _point_rows(item["points"], meal_name)]
        _sheet(wb, "메뉴별식수", [data["mode_name"], *POINT_HEADERS], rows, first=True)
        summary = [[item["name"], item["days"], item["avg_actual"], item["avg_usual"], item["avg_diff"]] for item in data["items"]]
        _sheet(wb, "메뉴별 요약", [data["mode_name"], "나온 날 수", "평균 실제(명)", "평균 평소(명)", "평균 차이(명)"], summary)
    elif tab == "popular":
        data = _run(analysis.popular_menus, db, meal_type, start, end, basis, order, limit, min_days, main_only)
        label = f"인기메뉴_{'상위' if order == 'top' else '하위'}{limit}"
        rows = [[item["rank"], item["name"], item["days"], item["avg_actual"], item["avg_usual"], item["avg_diff"]] for item in data["items"]]
        ws = _sheet(wb, "인기 메뉴", ["순위", data["basis_name"], "나온 날 수", "평균 실제(명)", "평균 평소(명)", "평균 차이(명)"], rows, first=True)
        ws.column_dimensions["B"].width = 28
        ws.append([])
        ws.append([f"기준: {data['basis_name']} · 최소 {data['min_days']}회 이상 · {'메인 메뉴로 나온 날만' if main_only else '모든 날'}"])
    elif tab == "ingredient":
        if ingredient_id is None:
            raise HTTPException(status_code=400, detail="재료를 선택해 주세요.")
        data = _run(analysis.ingredient, db, meal_type, start, end, ingredient_id, weather)
        label = f"재료별식수_{data['ingredient']['name']}"
        _sheet(wb, "재료별식수", POINT_HEADERS, _point_rows(data["points"], meal_name), first=True)
    elif tab == "weather":
        data = _run(analysis.weather_bands, db, meal_type, start, end)
        label = "날씨별식수"
        summary = [
            [group_name, band["label"], band["days"], band["avg_actual"], band["avg_usual"], band["avg_diff"], band["low_sample_note"] or ""]
            for group_name, groups in (("비", data["rain"]), ("기온", data["temperature"]))
            for band in groups
        ]
        _sheet(wb, "날씨별 요약", ["구분", "날씨", "일수", "평균 실제(명)", "평균 평소(명)", "평균 차이(명)", "참고"], summary, first=True)
        rows = []
        for group_name, groups in (("비", data["rain"]), ("기온", data["temperature"])):
            for band in groups:
                rows += [[f"{group_name}: {band['label']}", *row] for row in _point_rows(band["points"], meal_name)]
        _sheet(wb, "날짜 목록", ["날씨", *POINT_HEADERS], rows)
    else:
        raise HTTPException(status_code=400, detail="내려받을 화면이 올바르지 않습니다.")
    buffer = io.BytesIO()
    wb.save(buffer)
    filename = f"{label}_{meal_name}_{start.isoformat()}_{end.isoformat()}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"meal-count.xlsx\"; filename*=UTF-8''{quote(filename)}"},
    )
