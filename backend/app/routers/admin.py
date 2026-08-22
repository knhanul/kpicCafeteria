from __future__ import annotations

import hashlib
import shutil
import subprocess
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import settings
from ..db import get_db
from ..deps import admin_user
from ..models import (
    AuditLog,
    BackupRecord,
    DataArchive,
    Ingredient,
    IngredientAlias,
    MealActual,
    MealService,
    MealServiceMenu,
    MealServiceMenuIngredient,
    MealTypeSetting,
    Menu,
    PreservationRecord,
    Recipe,
    RecipeIngredient,
    User,
)

router = APIRouter(prefix="/api/admin", tags=["admin"])

ARCHIVE_RETENTION_HOURS = 24


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _parse_db_url(url: str) -> dict[str, str]:
    """Extract connection params from SQLAlchemy database URL for pg_dump."""
    # postgresql+psycopg://user:pass@host:port/dbname
    from urllib.parse import urlparse
    parsed = urlparse(url)
    return {
        "host": parsed.hostname or "localhost",
        "port": str(parsed.port or 5432),
        "user": parsed.username or "",
        "password": parsed.password or "",
        "dbname": parsed.path.lstrip("/"),
    }


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


# ==================== Backup ====================

@router.get("/backups")
def list_backups(db: Session = Depends(get_db), user: User = Depends(admin_user)):
    rows = db.scalars(
        select(BackupRecord).order_by(BackupRecord.created_at.desc())
    ).all()
    return {
        "items": [
            {
                "id": r.id,
                "filename": r.filename,
                "file_size": r.file_size,
                "backup_type": r.backup_type,
                "status": r.status,
                "created_at": r.created_at.isoformat() if r.created_at else None,
                "created_by": r.created_by,
            }
            for r in rows
        ]
    }


@router.post("/backups")
def create_backup(db: Session = Depends(get_db), user: User = Depends(admin_user)):
    if settings.database_url.startswith("sqlite"):
        raise HTTPException(status_code=400, detail="SQLite 환경에서는 백업을 지원하지 않습니다.")

    now = _utcnow()
    ts = now.strftime("%Y%m%d_%H%M%S")
    stored_filename = f"cafeteria_db_backup_manual_{ts}.dump"
    download_filename = f"cafeteria_db_backup_manual_{ts}.dump"
    dest = settings.backup_manual_dir / stored_filename

    params = _parse_db_url(settings.database_url)
    env = {
        "PGPASSWORD": params["password"],
        "PATH": __import__("os").environ.get("PATH", ""),
    }
    cmd = [
        "pg_dump",
        "-h", params["host"],
        "-p", params["port"],
        "-U", params["user"],
        "-d", params["dbname"],
        "-F", "c",
        "-f", str(dest),
    ]
    try:
        subprocess.run(cmd, env=env, check=True, capture_output=True, timeout=300)
    except FileNotFoundError:
        raise HTTPException(status_code=500, detail="백업 도구를 찾을 수 없습니다.")
    except subprocess.CalledProcessError:
        raise HTTPException(status_code=500, detail="시스템 데이터 백업에 실패했습니다.")
    except subprocess.TimeoutExpired:
        raise HTTPException(status_code=500, detail="시스템 데이터 백업에 실패했습니다.")

    file_size = dest.stat().st_size if dest.exists() else None
    if not file_size or file_size == 0:
        if dest.exists():
            dest.unlink()
        raise HTTPException(status_code=500, detail="시스템 데이터 백업에 실패했습니다.")

    checksum = _sha256(dest)
    record = BackupRecord(
        filename=download_filename,
        stored_filename=stored_filename,
        file_size=file_size,
        backup_type="manual",
        status="completed",
        checksum_sha256=checksum,
        created_by=user.username,
    )
    db.add(record)
    db.add(AuditLog(
        user_id=user.id,
        action="backup_create",
        entity_type="backup",
        entity_id=str(record.id) if record.id else None,
        detail={"filename": download_filename, "type": "manual"},
    ))
    db.commit()
    db.refresh(record)
    return {
        "id": record.id,
        "filename": record.filename,
        "file_size": record.file_size,
        "backup_type": record.backup_type,
        "status": record.status,
        "created_at": record.created_at.isoformat() if record.created_at else None,
    }


@router.get("/backups/{backup_id}/download")
def download_backup(backup_id: int, db: Session = Depends(get_db), user: User = Depends(admin_user)):
    record = db.get(BackupRecord, backup_id)
    if not record:
        raise HTTPException(status_code=404, detail="백업파일을 찾을 수 없습니다.")
    base_dir = settings.backup_dir.resolve()
    if record.backup_type == "auto":
        file_path = (settings.backup_auto_dir / record.stored_filename).resolve()
    else:
        file_path = (settings.backup_manual_dir / record.stored_filename).resolve()
    try:
        file_path.relative_to(base_dir)
    except ValueError:
        raise HTTPException(status_code=404, detail="백업파일을 찾을 수 없습니다.")
    if not file_path.is_file():
        raise HTTPException(status_code=404, detail="백업파일을 찾을 수 없습니다.")
    return FileResponse(
        path=str(file_path),
        media_type="application/octet-stream",
        filename=record.filename,
    )


@router.delete("/backups/{backup_id}")
def delete_backup(backup_id: int, db: Session = Depends(get_db), user: User = Depends(admin_user)):
    record = db.get(BackupRecord, backup_id)
    if not record:
        raise HTTPException(status_code=404, detail="백업파일을 찾을 수 없습니다.")
    base_dir = settings.backup_dir.resolve()
    if record.backup_type == "auto":
        file_path = (settings.backup_auto_dir / record.stored_filename).resolve()
    else:
        file_path = (settings.backup_manual_dir / record.stored_filename).resolve()
    try:
        file_path.relative_to(base_dir)
    except ValueError:
        raise HTTPException(status_code=404, detail="백업파일을 찾을 수 없습니다.")
    if file_path.is_file():
        file_path.unlink()
    db.add(AuditLog(
        user_id=user.id,
        action="backup_delete",
        entity_type="backup",
        entity_id=str(record.id),
        detail={"filename": record.filename},
    ))
    db.delete(record)
    db.commit()
    return {"ok": True}


# ==================== Excel Archive ====================

def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            val = cell.value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[letter].width = min(max_len + 4, 50)


def _add_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title=title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _auto_width(ws)
    ws.freeze_panes = "A2"


def _build_excel(db: Session, date_from: date | None, date_to: date | None) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    # 1. 식단 기록
    q = select(MealService).order_by(MealService.service_date)
    if date_from:
        q = q.where(MealService.service_date >= date_from)
    if date_to:
        q = q.where(MealService.service_date <= date_to)
    services = db.scalars(q).all()
    service_ids = [s.id for s in services]
    menus_map: dict[int, list[MealServiceMenu]] = {}
    if service_ids:
        menus = db.scalars(
            select(MealServiceMenu).where(MealServiceMenu.meal_service_id.in_(service_ids)).order_by(MealServiceMenu.sort_order)
        ).all()
        for m in menus:
            menus_map.setdefault(m.meal_service_id, []).append(m)
    _add_sheet(wb, "식단기록", [
        "날짜", "식사구분", "계획인원", "서비스시간", "콘셉트", "메뉴목록", "비고",
    ], [
        [
            s.service_date.isoformat(),
            s.meal_type,
            s.planned_count,
            s.service_time.isoformat() if s.service_time else "",
            s.concept_title or "",
            ", ".join(m.menu_name_snapshot for m in menus_map.get(s.id, [])),
            s.note or "",
        ]
        for s in services
    ])

    # 2. 조리지시서
    cooking_rows = []
    for s in services:
        for m in menus_map.get(s.id, []):
            cooking_rows.append([
                s.service_date.isoformat(),
                s.meal_type,
                m.menu_name_snapshot,
                m.note or "",
            ])
    _add_sheet(wb, "조리지시서", [
        "날짜", "식사구분", "메뉴명", "메뉴 비고",
    ], cooking_rows)

    # 3. 보존식 기록
    preservation_rows = []
    if service_ids:
        pres = db.scalars(
            select(PreservationRecord).where(PreservationRecord.meal_service_id.in_(service_ids))
        ).all()
        pres_map = {p.meal_service_id: p for p in pres}
        for s in services:
            p = pres_map.get(s.id)
            if p:
                preservation_rows.append([
                    s.service_date.isoformat(),
                    s.meal_type,
                    p.collection_time or "",
                    p.collector_name or "",
                    p.freezer_temperature or "",
                    p.manager_name or "",
                    p.disposal_at.isoformat() if p.disposal_at else "",
                    p.note or "",
                ])
    _add_sheet(wb, "보존식기록", [
        "날짜", "식사구분", "채수시간", "채수자", "냉동고온도", "관리자", "폐기시간", "비고",
    ], preservation_rows)

    # 4. 실제 식수
    actual_rows = []
    if service_ids:
        actuals = db.scalars(
            select(MealActual).where(MealActual.meal_service_id.in_(service_ids))
        ).all()
        act_map = {a.meal_service_id: a for a in actuals}
        for s in services:
            a = act_map.get(s.id)
            if a:
                actual_rows.append([
                    s.service_date.isoformat(),
                    s.meal_type,
                    a.actual_count or "",
                    a.note or "",
                    a.recorded_at.isoformat() if a.recorded_at else "",
                ])
    _add_sheet(wb, "실제식수", [
        "날짜", "식사구분", "실제인원", "비고", "기록시간",
    ], actual_rows)

    # 5. 메뉴 기준정보
    menus = db.scalars(select(Menu).order_by(Menu.name)).all()
    _add_sheet(wb, "메뉴기준정보", [
        "ID", "메뉴명", "역할", "사용여부", "검토상태",
    ], [
        [m.id, m.name, m.role, "사용" if m.active else "미사용", m.review_status]
        for m in menus
    ])

    # 6. 재료 기준정보
    ingredients = db.scalars(select(Ingredient).order_by(Ingredient.name)).all()
    _add_sheet(wb, "재료기준정보", [
        "ID", "재료명", "통계분석군", "기본단위", "kg환산계수", "사용여부",
    ], [
        [i.id, i.name, i.stat_group, i.default_unit or "", i.kg_factor or "", "사용" if i.active else "미사용"]
        for i in ingredients
    ])

    # 7. 레시피
    recipes = db.scalars(select(Recipe).order_by(Recipe.id)).all()
    recipe_rows = []
    for r in recipes:
        menu = db.get(Menu, r.menu_id)
        ri_rows = db.scalars(
            select(RecipeIngredient).where(RecipeIngredient.recipe_id == r.id).order_by(RecipeIngredient.sort_order)
        ).all()
        for ri in ri_rows:
            ing = db.get(Ingredient, ri.ingredient_id)
            recipe_rows.append([
                r.id,
                menu.name if menu else "",
                r.name,
                r.version,
                ing.name if ing else "",
                ri.quantity_per_100 or "",
                ri.unit or "",
            ])
    _add_sheet(wb, "레시피", [
        "레시피ID", "메뉴명", "레시피명", "버전", "재료명", "100인분량", "단위",
    ], recipe_rows)

    # 8. 식사유형 설정
    meal_types = db.scalars(select(MealTypeSetting).order_by(MealTypeSetting.sort_order)).all()
    _add_sheet(wb, "식사유형설정", [
        "코드", "이름", "기본인원", "기본시간", "정렬순서", "사용여부",
    ], [
        [mt.code, mt.name, mt.default_planned_count, mt.default_service_time.isoformat() if mt.default_service_time else "", mt.sort_order, "사용" if mt.active else "미사용"]
        for mt in meal_types
    ])

    # 9. 사용자 (password 제외)
    users = db.scalars(select(User).order_by(User.id)).all()
    _add_sheet(wb, "사용자목록", [
        "ID", "사용자ID", "사용자명", "권한", "사용여부", "최근로그인",
    ], [
        [u.id, u.username, u.display_name, u.role, "사용" if u.active else "사용중지", u.last_login_at.isoformat() if u.last_login_at else ""]
        for u in users
    ])

    return wb


@router.get("/archives")
def list_archives(db: Session = Depends(get_db), user: User = Depends(admin_user)):
    rows = db.scalars(
        select(DataArchive).order_by(DataArchive.created_at.desc())
    ).all()
    now = _utcnow()
    result = []
    for r in rows:
        expired = r.expires_at and r.expires_at < now
        file_path = settings.archive_dir / r.stored_filename
        file_exists = file_path.is_file()
        result.append({
            "id": r.id,
            "filename": r.filename,
            "file_size": r.file_size,
            "status": r.status if (file_exists and not expired) else "expired",
            "date_from": r.date_from.isoformat() if r.date_from else None,
            "date_to": r.date_to.isoformat() if r.date_to else None,
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "expires_at": r.expires_at.isoformat() if r.expires_at else None,
        })
    return {"items": result}


@router.post("/archives")
def create_archive(
    date_from: date | None = Query(None),
    date_to: date | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(admin_user),
):
    now = _utcnow()
    ts = now.strftime("%Y%m%d_%H%M%S")
    if date_from and date_to:
        label = f"{date_from.strftime('%Y%m%d')}_{date_to.strftime('%Y%m%d')}"
    else:
        label = "all"
    stored_filename = f"cafeteria_archive_{label}_{ts}.xlsx"
    download_filename = f"cafeteria_archive_{label}_{ts}.xlsx"
    dest = settings.archive_dir / stored_filename

    try:
        wb = _build_excel(db, date_from, date_to)
        wb.save(str(dest))
    except Exception:
        if dest.exists():
            dest.unlink()
        raise HTTPException(status_code=500, detail="Excel 데이터 아카이브 생성에 실패했습니다.")

    file_size = dest.stat().st_size if dest.exists() else None
    if not file_size or file_size == 0:
        if dest.exists():
            dest.unlink()
        raise HTTPException(status_code=500, detail="Excel 데이터 아카이브 생성에 실패했습니다.")

    expires_at = now + timedelta(hours=ARCHIVE_RETENTION_HOURS)
    record = DataArchive(
        filename=download_filename,
        stored_filename=stored_filename,
        file_size=file_size,
        status="completed",
        date_from=date_from,
        date_to=date_to,
        expires_at=expires_at,
    )
    db.add(record)
    db.add(AuditLog(
        user_id=user.id,
        action="archive_create",
        entity_type="data_archive",
        detail={"filename": download_filename},
    ))
    db.commit()
    db.refresh(record)
    return {
        "id": record.id,
        "filename": record.filename,
        "file_size": record.file_size,
        "status": record.status,
        "date_from": record.date_from.isoformat() if record.date_from else None,
        "date_to": record.date_to.isoformat() if record.date_to else None,
        "created_at": record.created_at.isoformat() if record.created_at else None,
        "expires_at": record.expires_at.isoformat() if record.expires_at else None,
    }


@router.get("/archives/{archive_id}/download")
def download_archive(archive_id: int, db: Session = Depends(get_db), user: User = Depends(admin_user)):
    record = db.get(DataArchive, archive_id)
    if not record:
        raise HTTPException(status_code=404, detail="아카이브 파일을 찾을 수 없습니다.")
    now = _utcnow()
    if record.expires_at and record.expires_at < now:
        raise HTTPException(status_code=404, detail="아카이브 파일을 찾을 수 없습니다.")
    base_dir = settings.archive_dir.resolve()
    file_path = (settings.archive_dir / record.stored_filename).resolve()
    try:
        file_path.relative_to(base_dir)
    except ValueError:
        raise HTTPException(status_code=404, detail="아카이브 파일을 찾을 수 없습니다.")
    if not file_path.is_file():
        raise HTTPException(status_code=404, detail="아카이브 파일을 찾을 수 없습니다.")
    return FileResponse(
        path=str(file_path),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=record.filename,
    )


@router.delete("/archives/{archive_id}")
def delete_archive(archive_id: int, db: Session = Depends(get_db), user: User = Depends(admin_user)):
    record = db.get(DataArchive, archive_id)
    if not record:
        raise HTTPException(status_code=404, detail="아카이브 파일을 찾을 수 없습니다.")
    base_dir = settings.archive_dir.resolve()
    file_path = (settings.archive_dir / record.stored_filename).resolve()
    try:
        file_path.relative_to(base_dir)
    except ValueError:
        raise HTTPException(status_code=404, detail="아카이브 파일을 찾을 수 없습니다.")
    if file_path.is_file():
        file_path.unlink()
    db.delete(record)
    db.commit()
    return {"ok": True}


@router.post("/archives/cleanup")
def cleanup_expired_archives(db: Session = Depends(get_db), user: User = Depends(admin_user)):
    now = _utcnow()
    archives = db.scalars(select(DataArchive).where(DataArchive.expires_at < now)).all()
    base_dir = settings.archive_dir.resolve()
    count = 0
    for a in archives:
        fp = (settings.archive_dir / a.stored_filename).resolve()
        try:
            fp.relative_to(base_dir)
            if fp.is_file():
                fp.unlink()
        except ValueError:
            pass
        db.delete(a)
        count += 1
    db.commit()
    return {"deleted": count}
