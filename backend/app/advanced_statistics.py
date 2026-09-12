from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, timedelta
from io import BytesIO
from math import sqrt
from statistics import median
from typing import Any

from openpyxl import Workbook
from sqlalchemy import and_, exists, func, or_, select
from sqlalchemy.orm import Session

from .models import MealActual, MealService, MealServiceMenu, Menu, WeatherHistory

WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]
MIN_SAMPLE = 5
TEMP_BUCKETS = ["<0", "0~4.9", "5~9.9", "10~14.9", "15~19.9", "20~24.9", "25~29.9", ">=30"]
HUMIDITY_BUCKETS = ["<0%", *[f"{value}~{value + 9.9:g}%" for value in range(0, 100, 10)], ">=100%"]


class StationSelectionError(ValueError):
    pass


def _meal_code(meal_type: str) -> str | None:
    return {"lunch": "LUNCH", "dinner": "DINNER"}.get(meal_type)


def _round(value: float | None, digits: int = 2) -> float | None:
    return round(value, digits) if value is not None else None


def _service_stmt(start: date, end: date, meal_type: str = "all"):
    stmt = (
        select(
            MealService.id,
            MealService.service_date,
            MealService.meal_type,
            MealService.planned_count,
            MealService.note.label("service_note"),
            MealActual.actual_count,
            MealActual.note.label("actual_note"),
        )
        .outerjoin(MealActual, MealActual.meal_service_id == MealService.id)
        .where(MealService.service_date.between(start, end))
    )
    code = _meal_code(meal_type)
    return stmt.where(MealService.meal_type == code) if code else stmt


def _service_rows(db: Session, start: date, end: date, meal_type: str = "all") -> list[dict[str, Any]]:
    stmt = _service_stmt(start, end, meal_type).order_by(MealService.service_date, MealService.meal_type, MealService.id)
    return [dict(row._mapping) for row in db.execute(stmt)]


def station_list(db: Session, start: date | None = None, end: date | None = None) -> dict[str, Any]:
    stmt = select(
        WeatherHistory.station_id,
        func.max(WeatherHistory.station_name).label("station_name"),
        func.min(WeatherHistory.observation_date).label("date_from"),
        func.max(WeatherHistory.observation_date).label("date_to"),
        func.count(WeatherHistory.id).label("record_count"),
    )
    if start:
        stmt = stmt.where(WeatherHistory.observation_date >= start)
    if end:
        stmt = stmt.where(WeatherHistory.observation_date <= end)
    rows = db.execute(stmt.group_by(WeatherHistory.station_id).order_by(WeatherHistory.station_id)).all()
    stations = [{
        "station_id": row.station_id,
        "station_name": row.station_name or "",
        "date_from": row.date_from.isoformat(),
        "date_to": row.date_to.isoformat(),
        "record_count": row.record_count,
    } for row in rows]
    return {"stations": stations, "auto_selected_station_id": stations[0]["station_id"] if len(stations) == 1 else None}


def _resolve_station(db: Session, station_id: str | None) -> tuple[str | None, str | None]:
    available = station_list(db)["stations"]
    if station_id:
        match = next((item for item in available if item["station_id"] == station_id), None)
        if not match:
            raise StationSelectionError("선택한 날씨 지점을 찾을 수 없습니다.")
        return match["station_id"], match["station_name"]
    if len(available) == 1:
        return available[0]["station_id"], available[0]["station_name"]
    if len(available) > 1:
        raise StationSelectionError("날씨 지점이 여러 개입니다. station_id를 지정해 주세요.")
    return None, None


def _weather_rows(db: Session, start: date, end: date, station_id: str | None) -> tuple[str | None, str | None, dict[date, dict[str, Any]]]:
    resolved_id, station_name = _resolve_station(db, station_id)
    if resolved_id is None:
        return None, None, {}
    rows = db.execute(
        select(
            WeatherHistory.observation_date,
            WeatherHistory.avg_temp,
            WeatherHistory.precipitation,
            WeatherHistory.avg_humidity,
            WeatherHistory.snow_depth,
            WeatherHistory.sunshine_hours,
        ).where(
            WeatherHistory.observation_date.between(start, end),
            WeatherHistory.station_id == resolved_id,
        )
    )
    return resolved_id, station_name, {row.observation_date: dict(row._mapping) for row in rows}


def _optional_weather_rows(db: Session, start: date, end: date, station_id: str | None) -> tuple[str | None, str | None, dict[date, dict[str, Any]]]:
    if station_id:
        return _weather_rows(db, start, end, station_id)
    stations = station_list(db)["stations"]
    if len(stations) == 1:
        return _weather_rows(db, start, end, stations[0]["station_id"])
    return None, None, {}


def _menu_rows(db: Session, service_ids: list[int]) -> list[dict[str, Any]]:
    if not service_ids:
        return []
    rows = db.execute(
        select(
            MealServiceMenu.meal_service_id,
            MealServiceMenu.menu_id,
            Menu.canonical_name,
            Menu.role,
            MealServiceMenu.menu_name_snapshot,
        )
        .outerjoin(Menu, Menu.id == MealServiceMenu.menu_id)
        .where(MealServiceMenu.meal_service_id.in_(service_ids), MealServiceMenu.is_representative.is_(True))
        .order_by(MealServiceMenu.meal_service_id, MealServiceMenu.sort_order, MealServiceMenu.id)
    )
    result = []
    seen: set[tuple[int, str]] = set()
    for row in rows:
        linked_name = (row.canonical_name or "").strip()
        canonical = linked_name or (row.menu_name_snapshot or "").strip()
        if not canonical or (row.meal_service_id, canonical) in seen:
            continue
        seen.add((row.meal_service_id, canonical))
        result.append({
            "meal_service_id": row.meal_service_id,
            "menu_id": row.menu_id,
            "canonical_name": canonical,
            "role": row.role or "기타",
            "canonical_linked": bool(linked_name),
        })
    return result


def _sample(values: list[int | float]) -> dict[str, Any]:
    return {
        "average": _round(sum(values) / len(values)) if values else None,
        "median": _round(float(median(values))) if values else None,
        "minimum": min(values) if values else None,
        "maximum": max(values) if values else None,
        "n": len(values),
    }


def _service_sample(rows: list[dict[str, Any]]) -> dict[str, Any]:
    result = _sample([row["actual_count"] for row in rows])
    result["average_planned"] = _round(sum(row["planned_count"] for row in rows) / len(rows)) if rows else None
    result["average_plan_error"] = _round(sum(row["planned_count"] - row["actual_count"] for row in rows) / len(rows)) if rows else None
    return result


def _comparison(rows: list[dict[str, Any]]) -> dict[str, Any]:
    actual_rows = [row for row in rows if row["actual_count"] is not None]
    actual_total = sum(row["actual_count"] for row in actual_rows)
    comparison_planned = sum(row["planned_count"] for row in actual_rows)
    errors = [row["planned_count"] - row["actual_count"] for row in actual_rows]
    eligible = [(row, error) for row, error in zip(actual_rows, errors) if row["actual_count"] != 0]
    within5 = sum(abs(error) / abs(row["actual_count"]) <= .05 for row, error in eligible)
    within10 = sum(abs(error) / abs(row["actual_count"]) <= .10 for row, error in eligible)
    return {
        "n": len(actual_rows),
        "planned_sum": comparison_planned,
        "comparison_planned_sum": comparison_planned,
        "actual_sum": actual_total if actual_rows else None,
        "plan_minus_actual": comparison_planned - actual_total if actual_rows else None,
        "mae": _round(sum(abs(value) for value in errors) / len(errors)) if errors else None,
        "wape": _round(sum(abs(value) for value in errors) / actual_total * 100) if actual_total != 0 else None,
        "bias": _round(sum(errors) / len(errors)) if errors else None,
        "bias_rate": _round(sum(errors) / actual_total * 100) if actual_total != 0 else None,
        "actual_over_plan_count": sum(value < 0 for value in errors),
        "actual_under_plan_count": sum(value > 0 for value in errors),
        "over_count": sum(value > 0 for value in errors),
        "under_count": sum(value < 0 for value in errors),
        "exact_count": sum(value == 0 for value in errors),
        "threshold_n": len(eligible),
        "within_5_count": within5,
        "within_5_rate": _round(within5 / len(eligible) * 100) if eligible else None,
        "within_10_count": within10,
        "within_10_rate": _round(within10 / len(eligible) * 100) if eligible else None,
    }


def _series(rows: list[dict[str, Any]], key_func, labels: list[Any] | None = None) -> list[dict[str, Any]]:
    grouped: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[key_func(row)].append(row)
    keys = labels if labels is not None else sorted(grouped)
    result = []
    for key in keys:
        subset = grouped.get(key, [])
        comparison = _comparison(subset)
        result.append({
            "key": key,
            "service_count": len(subset),
            "planned_sum": sum(row["planned_count"] for row in subset),
            "analyzed_count": comparison["n"],
            "comparison_planned_sum": comparison["comparison_planned_sum"],
            "planned_average": _round(sum(row["planned_count"] for row in subset) / len(subset)) if subset else None,
            "actual_sum": comparison["actual_sum"],
            "actual_average": _round(comparison["actual_sum"] / comparison["n"]) if comparison["n"] else None,
        })
    return result


def _insights(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    actual_rows = [row for row in rows if row["actual_count"] is not None]
    if len(actual_rows) < MIN_SAMPLE:
        return [{"code": "insufficient_sample", "n": len(actual_rows), "message": "비교 가능한 실제식수 표본이 5건 미만입니다."}]
    comparison = _comparison(actual_rows)
    weekday_groups = defaultdict(list)
    for row in actual_rows:
        weekday_groups[row["service_date"].weekday()].append(row["actual_count"])
    eligible = [(weekday, values) for weekday, values in weekday_groups.items() if len(values) >= MIN_SAMPLE]
    bias_rate_text = f"{comparison['bias_rate']:+.2f}%" if comparison["bias_rate"] is not None else "산출 불가"
    insights = [{
        "code": "plan_bias",
        "n": comparison["n"],
        "message": f"계획-실제 평균 오차는 {comparison['bias']:+.2f}식, Bias율은 {bias_rate_text}입니다.",
    }]
    if eligible:
        weekday, values = max(eligible, key=lambda item: (sum(item[1]) / len(item[1]), -item[0]))
        insights.append({
            "code": "highest_weekday",
            "n": len(values),
            "message": f"표본 5건 이상 요일 중 {WEEKDAYS[weekday]}요일 실제식수 평균이 가장 높습니다.",
        })
    missing = len(rows) - len(actual_rows)
    if missing:
        insights.append({"code": "missing_actual", "n": missing, "message": f"실제식수 미입력 서비스가 {missing}건입니다."})
    return insights


def overview(db: Session, start: date, end: date, meal_type: str = "all") -> dict[str, Any]:
    rows = _service_rows(db, start, end, meal_type)
    comparison = _comparison(rows)
    daily = _series(rows, lambda row: row["service_date"].isoformat())
    weekly = _series(rows, lambda row: (row["service_date"] - timedelta(days=row["service_date"].weekday())).isoformat())
    monthly = _series(rows, lambda row: row["service_date"].strftime("%Y-%m"))
    weekday = _series(rows, lambda row: row["service_date"].weekday(), list(range(7)))
    for item in weekday:
        item["weekday"] = WEEKDAYS[item.pop("key")]
    meal_types = _series(rows, lambda row: row["meal_type"], ["LUNCH", "DINNER"])
    return {
        "date_from": start.isoformat(),
        "date_to": end.isoformat(),
        "meal_type": meal_type,
        "service_count": len(rows),
        "analyzed_count": comparison["n"],
        "missing_actual_count": len(rows) - comparison["n"],
        "actual_zero_count": sum(row["actual_count"] == 0 for row in rows if row["actual_count"] is not None),
        "actual_input_rate": _round(comparison["n"] / len(rows) * 100) if rows else None,
        "planned_sum": sum(row["planned_count"] for row in rows),
        "comparison_planned_sum": comparison["comparison_planned_sum"],
        "actual_sum": comparison["actual_sum"],
        "difference": comparison["actual_sum"] - comparison["comparison_planned_sum"] if comparison["actual_sum"] is not None else None,
        "actual": _sample([row["actual_count"] for row in rows if row["actual_count"] is not None]),
        "daily": daily,
        "weekly": weekly,
        "monthly": monthly,
        "weekday": weekday,
        "meal_types": meal_types,
        "insights": _insights(rows),
    }


def _error_distribution(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    labels = ["<=-20%", "-20~-10%", "-10~-5%", "-5~5%", "5~10%", "10~20%", ">=20%", "actual=0"]
    counts = {label: 0 for label in labels}
    for row in rows:
        actual = row["actual_count"]
        if actual == 0:
            counts["actual=0"] += 1
            continue
        rate = (row["planned_count"] - actual) / abs(actual) * 100
        label = "<=-20%" if rate <= -20 else "-20~-10%" if rate < -10 else "-10~-5%" if rate < -5 else "-5~5%" if rate <= 5 else "5~10%" if rate <= 10 else "10~20%" if rate < 20 else ">=20%"
        counts[label] += 1
    return [{"bucket": label, "count": counts[label]} for label in labels]


def _percentile(values: list[float], fraction: float) -> float:
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def _outliers(rows: list[dict[str, Any]], menus: list[dict[str, Any]], weather: dict[date, dict[str, Any]]) -> list[dict[str, Any]]:
    menus_by_service: dict[int, list[str]] = defaultdict(list)
    for item in menus:
        menus_by_service[item["meal_service_id"]].append(item["canonical_name"])
    result = []
    for meal_type in ("LUNCH", "DINNER"):
        subset = [row for row in rows if row["meal_type"] == meal_type and row["actual_count"] is not None]
        if len(subset) < 4:
            continue
        values = [float(row["actual_count"]) for row in subset]
        q1, q3 = _percentile(values, .25), _percentile(values, .75)
        iqr = q3 - q1
        low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        for row in subset:
            if row["actual_count"] < low or row["actual_count"] > high:
                result.append({
                    "meal_service_id": row["id"],
                    "service_date": row["service_date"].isoformat(),
                    "meal_type": meal_type,
                    "actual_count": row["actual_count"],
                    "planned_count": row["planned_count"],
                    "lower_fence": _round(low),
                    "upper_fence": _round(high),
                    "representative_menus": menus_by_service.get(row["id"], []),
                    "weather": weather.get(row["service_date"]),
                    "note": row["actual_note"] or row["service_note"],
                })
    return result


def plan_vs_actual(db: Session, start: date, end: date, meal_type: str = "all", station_id: str | None = None) -> dict[str, Any]:
    all_rows = _service_rows(db, start, end, meal_type)
    rows = [row for row in all_rows if row["actual_count"] is not None]
    result = _comparison(rows)
    daily_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    weekly_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    monthly_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    weekday_groups: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        daily_groups[row["service_date"].isoformat()].append(row)
        weekly_groups[(row["service_date"] - timedelta(days=row["service_date"].weekday())).isoformat()].append(row)
        monthly_groups[row["service_date"].strftime("%Y-%m")].append(row)
        weekday_groups[row["service_date"].weekday()].append(row)
    result["daily"] = [{"period": key, **_comparison(daily_groups[key])} for key in sorted(daily_groups)]
    result["weekly"] = [{"period": key, **_comparison(weekly_groups[key])} for key in sorted(weekly_groups)]
    result["monthly"] = [{"period": key, "month": key, **_comparison(monthly_groups[key])} for key in sorted(monthly_groups)]
    result["weekday"] = [{
        "weekday": WEEKDAYS[index],
        "n": len(weekday_groups[index]),
        "average_error": _round(sum(row["planned_count"] - row["actual_count"] for row in weekday_groups[index]) / len(weekday_groups[index])) if weekday_groups[index] else None,
    } for index in range(7)]
    result["error_distribution"] = _error_distribution(rows)
    result["error_definition"] = "planned_count - actual_count; percentage denominators use non-zero actual_count"
    resolved_id, station_name, weather = _optional_weather_rows(db, start, end, station_id)
    menus = _menu_rows(db, [row["id"] for row in rows])
    result["station_id"] = resolved_id
    result["station_name"] = station_name
    result["outlier_method"] = "meal_type별 actual_count Q1-1.5×IQR 미만 또는 Q3+1.5×IQR 초과"
    result["outliers"] = _outliers(rows, menus, weather)
    result["insights"] = _insights(all_rows)
    return result


def menu_metrics(db: Session, start: date, end: date, meal_type: str = "all") -> dict[str, Any]:
    services = [row for row in _service_rows(db, start, end, meal_type) if row["actual_count"] is not None]
    service_by_id = {row["id"]: row for row in services}
    links = _menu_rows(db, list(service_by_id))
    menu_services: dict[str, list[dict[str, Any]]] = defaultdict(list)
    menu_ids: dict[str, set[int]] = defaultdict(set)
    role_services: dict[str, dict[int, dict[str, Any]]] = defaultdict(dict)
    service_menus: dict[int, list[str]] = defaultdict(list)
    for link in links:
        service = service_by_id.get(link["meal_service_id"])
        if service:
            menu_services[link["canonical_name"]].append(service)
            role_services[link["role"]][service["id"]] = service
            service_menus[service["id"]].append(link["canonical_name"])
            if link["menu_id"] is not None:
                menu_ids[link["canonical_name"]].add(link["menu_id"])
    strata: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for service in services:
        strata[(service["meal_type"], service["service_date"].weekday())].append(service)
    items = []
    for name, occurrences in menu_services.items():
        occurrence_ids = {row["id"] for row in occurrences}
        comparator: dict[int, dict[str, Any]] = {}
        matched_baselines: list[float] = []
        for occurrence in occurrences:
            peers = [peer for peer in strata[(occurrence["meal_type"], occurrence["service_date"].weekday())] if peer["id"] not in occurrence_ids]
            for peer in peers:
                comparator[peer["id"]] = peer
            if peers:
                matched_baselines.append(sum(peer["actual_count"] for peer in peers) / len(peers))
        values = [row["actual_count"] for row in occurrences]
        planned = [row["planned_count"] for row in occurrences]
        peers = list(comparator.values())
        eligible = len(occurrences) >= MIN_SAMPLE and len(peers) >= MIN_SAMPLE and len(matched_baselines) == len(occurrences)
        menu_average = sum(values) / len(values)
        comparator_average = sum(matched_baselines) / len(matched_baselines) if matched_baselines else None
        items.append({
            "canonical_name": name,
            "menu_ids": sorted(menu_ids[name]),
            **_sample(values),
            "occurrence_n": len(occurrences),
            "average_planned": _round(sum(planned) / len(planned)),
            "average_plan_error": _round(sum(p - a for p, a in zip(planned, values)) / len(values)),
            "comparator_n": len(peers),
            "comparator_average": _round(comparator_average),
            "lift": _round(menu_average - comparator_average) if eligible else None,
            "lift_percent": _round((menu_average - comparator_average) / comparator_average * 100) if eligible and comparator_average != 0 else None,
            "lift_eligible": eligible,
            "sample_ok": len(occurrences) >= MIN_SAMPLE,
        })
    items.sort(key=lambda item: (-item["occurrence_n"], item["canonical_name"]))
    eligible_items = [item for item in items if item["lift_eligible"]]
    ranking_count = min(10, len(items))
    rankings = {
        "highest_actual": sorted(items, key=lambda item: (-(item["average"] or 0), item["canonical_name"]))[:ranking_count],
        "lowest_actual": sorted(items, key=lambda item: ((item["average"] or 0), item["canonical_name"]))[:ranking_count],
        "largest_over_plan": sorted(items, key=lambda item: (-(item["average_plan_error"] or 0), item["canonical_name"]))[:ranking_count],
        "largest_under_plan": sorted(items, key=lambda item: ((item["average_plan_error"] or 0), item["canonical_name"]))[:ranking_count],
        "highest_lift": sorted(eligible_items, key=lambda item: (-item["lift"], item["canonical_name"]))[:10],
        "lowest_lift": sorted(eligible_items, key=lambda item: (item["lift"], item["canonical_name"]))[:10],
    }
    sample_items = [item for item in items if item["sample_ok"]]
    insights = []
    if sample_items:
        highest = max(sample_items, key=lambda item: item["average"])
        insights.append({"code": "highest_menu", "n": highest["n"], "message": f"제공 5회 이상 대표메뉴 중 {highest['canonical_name']} 제공일의 평균 실제식수가 가장 높았습니다."})
    else:
        insights.append({"code": "insufficient_sample", "n": 0, "message": "제공 5회 이상인 대표메뉴가 없어 메뉴 평균 비교를 수행하지 않았습니다."})
    if rankings["highest_lift"]:
        lifted = rankings["highest_lift"][0]
        insights.append({"code": "highest_lift", "n": lifted["n"], "comparator_n": lifted["comparator_n"], "message": f"{lifted['canonical_name']} 제공일은 동일 요일·배식유형 비교군보다 평균 {lifted['lift_percent']:+.1f}% 높게 나타났습니다."})
    role_items = [{"role": role, **_service_sample(list(by_service.values()))} for role, by_service in sorted(role_services.items())]
    combinations = Counter(tuple(sorted(set(names))) for names in service_menus.values() if names)
    combination_items = [{"menus": list(names), "count": count} for names, count in combinations.most_common(20)]
    return {"items": items, "rankings": rankings, "roles": role_items, "combinations": combination_items, "menu_count": len(items), "service_count": len(services), "minimum_sample": MIN_SAMPLE, "insights": insights}


def _season(value: date) -> str:
    if value.month in (3, 4, 5):
        return "봄"
    if value.month in (6, 7, 8):
        return "여름"
    if value.month in (9, 10, 11):
        return "가을"
    return "겨울"


def _temp_bucket(value: float) -> str:
    if value < 0:
        return "<0"
    if value < 5:
        return "0~4.9"
    if value < 10:
        return "5~9.9"
    if value < 15:
        return "10~14.9"
    if value < 20:
        return "15~19.9"
    if value < 25:
        return "20~24.9"
    if value < 30:
        return "25~29.9"
    return ">=30"


def _humidity_bucket(value: float) -> str:
    if value < 0:
        return "<0%"
    if value >= 100:
        return ">=100%"
    lower = int(value // 10) * 10
    return f"{lower}~{lower + 9.9:g}%"


def _pearson(pairs: list[tuple[float, float]]) -> float | None:
    if len(pairs) < 10:
        return None
    xs, ys = zip(*pairs)
    x_mean, y_mean = sum(xs) / len(xs), sum(ys) / len(ys)
    numerator = sum((x - x_mean) * (y - y_mean) for x, y in pairs)
    denominator = sqrt(sum((x - x_mean) ** 2 for x in xs) * sum((y - y_mean) ** 2 for y in ys))
    return _round(numerator / denominator, 4) if denominator else None


def _group_samples(groups: dict[str, list[int]], order: list[str] | None = None) -> list[dict[str, Any]]:
    keys = order or sorted(groups)
    return [{"bucket": key, **_sample(groups[key])} for key in keys]


def _group_service_samples(groups: dict[str, list[dict[str, Any]]], order: list[str]) -> list[dict[str, Any]]:
    return [{"bucket": key, **_service_sample(groups[key])} for key in order]


def weather_metrics(db: Session, start: date, end: date, meal_type: str = "all", station_id: str | None = None) -> dict[str, Any]:
    services = [row for row in _service_rows(db, start, end, meal_type) if row["actual_count"] is not None]
    resolved_id, station_name, weather = _weather_rows(db, start, end, station_id)
    temperature: dict[str, list[dict[str, Any]]] = defaultdict(list)
    humidity: dict[str, list[dict[str, Any]]] = defaultdict(list)
    rain: dict[str, list[dict[str, Any]]] = defaultdict(list)
    snow: dict[str, list[dict[str, Any]]] = defaultdict(list)
    seasons: dict[str, list[dict[str, Any]]] = defaultdict(list)
    missing = {"weather": 0, "avg_temp": 0, "precipitation": 0, "avg_humidity": 0, "snow_depth": 0}
    scatter = []
    for service in services:
        actual = service["actual_count"]
        seasons[_season(service["service_date"])].append(service)
        weather_row = weather.get(service["service_date"])
        if not weather_row:
            missing["weather"] += 1
            continue
        for field in ("avg_temp", "precipitation", "avg_humidity", "snow_depth"):
            if weather_row[field] is None:
                missing[field] += 1
        if weather_row["avg_temp"] is not None:
            temperature[_temp_bucket(weather_row["avg_temp"])].append(service)
        if weather_row["avg_humidity"] is not None:
            humidity[_humidity_bucket(weather_row["avg_humidity"])].append(service)
        if weather_row["precipitation"] is None:
            rain["NULL"].append(service)
        elif weather_row["precipitation"] == 0:
            rain["0"].append(service)
        else:
            rain[">0"].append(service)
        if weather_row["snow_depth"] is not None:
            snow["0" if weather_row["snow_depth"] == 0 else ">0"].append(service)
        scatter.append({"service_date": service["service_date"].isoformat(), "meal_type": service["meal_type"], "actual_count": actual, **weather_row})
    dry_average = sum(row["actual_count"] for row in rain["0"]) / len(rain["0"]) if rain["0"] else None
    rain_summary = []
    for bucket in ("NULL", "0", ">0"):
        values = rain[bucket]
        average = sum(row["actual_count"] for row in values) / len(values) if values else None
        difference = average - dry_average if average is not None and dry_average is not None else None
        rain_summary.append({
            "bucket": bucket,
            **_service_sample(values),
            "difference_vs_zero": _round(difference),
            "percent_vs_zero": _round(difference / dry_average * 100) if difference is not None and dry_average != 0 else None,
        })
    correlations = {}
    for field in ("avg_temp", "precipitation", "avg_humidity", "snow_depth", "sunshine_hours"):
        pairs = [(float(weather[service["service_date"]][field]), float(service["actual_count"])) for service in services if service["service_date"] in weather and weather[service["service_date"]][field] is not None]
        correlations[field] = {"pearson": _pearson(pairs), "n": len(pairs), "eligible": len(pairs) >= 10, "minimum_n": 10}
    insights = []
    dry = next(item for item in rain_summary if item["bucket"] == "0")
    wet = next(item for item in rain_summary if item["bucket"] == ">0")
    if dry["n"] >= MIN_SAMPLE and wet["n"] >= MIN_SAMPLE and wet["percent_vs_zero"] is not None:
        insights.append({"code": "rain_difference", "n": wet["n"], "comparison_n": dry["n"], "message": f"강수 기록이 있는 날의 평균 실제식수는 무강수일보다 {wet['percent_vs_zero']:+.1f}% 차이가 나타났습니다."})
    else:
        insights.append({"code": "rain_insufficient", "n": wet["n"], "comparison_n": dry["n"], "message": "강수일 또는 무강수일 표본이 5건 미만이어서 평균 차이를 해석하지 않았습니다."})
    if correlations["avg_temp"]["eligible"]:
        insights.append({"code": "temperature_correlation", "n": correlations["avg_temp"]["n"], "message": f"평균기온과 실제식수의 Pearson 상관계수는 {correlations['avg_temp']['pearson']:+.3f}입니다. 상관관계이며 인과관계를 의미하지 않습니다."})
    return {
        "station_id": resolved_id,
        "station_name": station_name,
        "matched_service_count": len(services) - missing["weather"],
        "missing_weather_service_count": missing["weather"],
        "missing_counts": missing,
        "temperature": _group_service_samples(temperature, TEMP_BUCKETS),
        "humidity": _group_service_samples(humidity, HUMIDITY_BUCKETS),
        "rain": rain_summary,
        "snow": {"meaningful_n": sum(len(values) for values in snow.values()), "missing_n": missing["weather"] + missing["snow_depth"], "items": _group_service_samples(snow, ["0", ">0"])},
        "seasons": _group_service_samples(seasons, ["봄", "여름", "가을", "겨울"]),
        "scatter": scatter,
        "correlations": correlations,
        "insights": insights,
        "correlation_disclaimer": "Pearson 상관계수는 유효한 짝 표본 N>=10인 경우에만 제공되며 인과관계를 의미하지 않습니다.",
    }


def menu_weather_metrics(db: Session, start: date, end: date, meal_type: str = "all", station_id: str | None = None) -> dict[str, Any]:
    services = [row for row in _service_rows(db, start, end, meal_type) if row["actual_count"] is not None]
    service_by_id = {row["id"]: row for row in services}
    menus = _menu_rows(db, list(service_by_id))
    resolved_id, station_name, weather = _weather_rows(db, start, end, station_id)
    weekday_temp: dict[tuple[int, str], list[int]] = defaultdict(list)
    weekday_rain: dict[tuple[int, str], list[int]] = defaultdict(list)
    menu_temp: dict[tuple[str, str], list[int]] = defaultdict(list)
    menu_rain: dict[tuple[str, str], list[int]] = defaultdict(list)
    menus_by_service: dict[int, list[str]] = defaultdict(list)
    for menu in menus:
        menus_by_service[menu["meal_service_id"]].append(menu["canonical_name"])
    for service in services:
        weather_row = weather.get(service["service_date"])
        if not weather_row:
            continue
        weekday = service["service_date"].weekday()
        temp_bucket = _temp_bucket(weather_row["avg_temp"]) if weather_row["avg_temp"] is not None else None
        rain_bucket = None if weather_row["precipitation"] is None else "0" if weather_row["precipitation"] == 0 else ">0"
        if temp_bucket:
            weekday_temp[(weekday, temp_bucket)].append(service["actual_count"])
        if rain_bucket:
            weekday_rain[(weekday, rain_bucket)].append(service["actual_count"])
        for menu_name in menus_by_service.get(service["id"], []):
            if temp_bucket:
                menu_temp[(menu_name, temp_bucket)].append(service["actual_count"])
            if rain_bucket:
                menu_rain[(menu_name, rain_bucket)].append(service["actual_count"])
    def heatmap(groups, menu_dimension=False):
        result = []
        for key, values in groups.items():
            item = {"bucket": key[1], **_sample(values), "sample_ok": len(values) >= MIN_SAMPLE}
            item["canonical_name" if menu_dimension else "weekday"] = key[0] if menu_dimension else WEEKDAYS[key[0]]
            result.append(item)
        return sorted(result, key=lambda item: (item.get("canonical_name", item.get("weekday")), item["bucket"]))
    return {
        "station_id": resolved_id,
        "station_name": station_name,
        "minimum_sample": MIN_SAMPLE,
        "weekday_temperature": heatmap(weekday_temp),
        "weekday_rain": heatmap(weekday_rain),
        "representative_menu_temperature": heatmap(menu_temp, True),
        "representative_menu_rain": heatmap(menu_rain, True),
    }


def _quality(count: int, total: int) -> dict[str, Any]:
    return {"count": count, "rate": _round(count / total * 100) if total else None, "denominator": total}


def data_quality(db: Session, start: date, end: date, meal_type: str = "all", station_id: str | None = None) -> dict[str, Any]:
    services = _service_rows(db, start, end, meal_type)
    service_ids = [row["id"] for row in services]
    menus = _menu_rows(db, service_ids)
    services_with_menu = {row["meal_service_id"] for row in menus}
    canonical_menu_count = sum(row["canonical_linked"] for row in menus)
    resolved_id, station_name, weather = _weather_rows(db, start, end, station_id)
    total = len(services)
    actual_count = sum(row["actual_count"] is not None for row in services)
    weather_count = sum(row["service_date"] in weather for row in services)
    temp_count = sum(row["service_date"] in weather and weather[row["service_date"]]["avg_temp"] is not None for row in services)
    rain_count = sum(row["service_date"] in weather and weather[row["service_date"]]["precipitation"] is not None for row in services)
    metrics = {
        "actual": _quality(actual_count, total),
        "planned": _quality(sum(row["planned_count"] is not None for row in services), total),
        "representative_menu": _quality(len(services_with_menu), total),
        "canonical_linkage": _quality(canonical_menu_count, len(menus)),
        "weather_match": _quality(weather_count, total),
        "avg_temp": _quality(temp_count, total),
        "precipitation": _quality(rain_count, total),
    }
    return {
        "service_count": total,
        "station_id": resolved_id,
        "station_name": station_name,
        "metrics": metrics,
        "actual_present_count": actual_count,
        "actual_missing_count": total - actual_count,
        "actual_zero_count": sum(row["actual_count"] == 0 for row in services if row["actual_count"] is not None),
        "planned_present_count": metrics["planned"]["count"],
        "representative_menu_present_count": len(services_with_menu),
        "representative_menu_missing_count": total - len(services_with_menu),
        "canonical_linkage_count": canonical_menu_count,
        "weather_matched_service_count": weather_count,
        "weather_missing_service_count": total - weather_count,
    }


def _filtered_detail_stmt(start: date, end: date, meal_type: str, resolved_id: str | None, rain: str, menu: str | None, temp_bucket: str | None):
    stmt = _service_stmt(start, end, meal_type)
    if resolved_id:
        stmt = stmt.outerjoin(
            WeatherHistory,
            and_(WeatherHistory.observation_date == MealService.service_date, WeatherHistory.station_id == resolved_id),
        ).add_columns(
            WeatherHistory.avg_temp,
            WeatherHistory.precipitation,
            WeatherHistory.avg_humidity,
            WeatherHistory.snow_depth,
            WeatherHistory.sunshine_hours,
        )
        if rain == "missing":
            stmt = stmt.where(or_(WeatherHistory.id.is_(None), WeatherHistory.precipitation.is_(None)))
        elif rain == "dry":
            stmt = stmt.where(WeatherHistory.precipitation == 0)
        elif rain == "rain":
            stmt = stmt.where(WeatherHistory.precipitation > 0)
        if temp_bucket:
            ranges = {
                "<0": WeatherHistory.avg_temp < 0,
                "0~4.9": and_(WeatherHistory.avg_temp >= 0, WeatherHistory.avg_temp < 5),
                "5~9.9": and_(WeatherHistory.avg_temp >= 5, WeatherHistory.avg_temp < 10),
                "10~14.9": and_(WeatherHistory.avg_temp >= 10, WeatherHistory.avg_temp < 15),
                "15~19.9": and_(WeatherHistory.avg_temp >= 15, WeatherHistory.avg_temp < 20),
                "20~24.9": and_(WeatherHistory.avg_temp >= 20, WeatherHistory.avg_temp < 25),
                "25~29.9": and_(WeatherHistory.avg_temp >= 25, WeatherHistory.avg_temp < 30),
                ">=30": WeatherHistory.avg_temp >= 30,
            }
            stmt = stmt.where(ranges[temp_bucket])
    if menu:
        term = f"%{menu.strip()}%"
        menu_match = exists(select(MealServiceMenu.id).outerjoin(Menu, Menu.id == MealServiceMenu.menu_id).where(
            MealServiceMenu.meal_service_id == MealService.id,
            MealServiceMenu.is_representative.is_(True),
            or_(Menu.canonical_name.ilike(term), MealServiceMenu.menu_name_snapshot.ilike(term)),
        ))
        stmt = stmt.where(menu_match)
    return stmt


def _detail_items(db: Session, stmt, resolved_id: str | None, offset: int | None = None, limit: int | None = None) -> list[dict[str, Any]]:
    ordered = stmt.order_by(MealService.service_date, MealService.meal_type, MealService.id)
    if offset is not None:
        ordered = ordered.offset(offset)
    if limit is not None:
        ordered = ordered.limit(limit)
    rows = [dict(row._mapping) for row in db.execute(ordered)]
    menus = _menu_rows(db, [row["id"] for row in rows])
    menus_by_service: dict[int, list[str]] = defaultdict(list)
    for item in menus:
        menus_by_service[item["meal_service_id"]].append(item["canonical_name"])
    items = []
    for row in rows:
        weather = None
        if resolved_id:
            weather = {field: row.get(field) for field in ("avg_temp", "precipitation", "avg_humidity", "snow_depth", "sunshine_hours")}
        items.append({
            "meal_service_id": row["id"],
            "service_date": row["service_date"].isoformat(),
            "meal_type": row["meal_type"],
            "planned_count": row["planned_count"],
            "actual_count": row["actual_count"],
            "plan_error": row["planned_count"] - row["actual_count"] if row["actual_count"] is not None else None,
            "representative_menus": menus_by_service.get(row["id"], []),
            "weather": weather,
            "note": row["actual_note"] or row["service_note"],
        })
    return items


def drilldown(db: Session, start: date, end: date, meal_type: str = "all", station_id: str | None = None, page: int = 1, page_size: int = 50, rain: str = "all", menu: str | None = None, temp_bucket: str | None = None) -> dict[str, Any]:
    resolved_id, station_name = _resolve_station(db, station_id)
    if (rain != "all" or temp_bucket) and not resolved_id:
        raise StationSelectionError("날씨 상세 필터를 사용하려면 날씨 지점이 필요합니다.")
    stmt = _filtered_detail_stmt(start, end, meal_type, resolved_id, rain, menu, temp_bucket)
    total = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery())) or 0
    items = _detail_items(db, stmt, resolved_id, (page - 1) * page_size, page_size)
    return {"total": total, "page": page, "page_size": page_size, "station_id": resolved_id, "station_name": station_name, "filters": {"rain": rain, "menu": menu, "temp_bucket": temp_bucket}, "items": items}


def export_xlsx(db: Session, start: date, end: date, meal_type: str = "all", station_id: str | None = None, rain: str = "all", menu: str | None = None, temp_bucket: str | None = None) -> bytes:
    resolved_id, station_name = _resolve_station(db, station_id)
    if (rain != "all" or temp_bucket) and not resolved_id:
        raise StationSelectionError("날씨 상세 필터를 사용하려면 날씨 지점이 필요합니다.")
    stmt = _filtered_detail_stmt(start, end, meal_type, resolved_id, rain, menu, temp_bucket)
    items = _detail_items(db, stmt, resolved_id)
    comparable = [item for item in items if item["actual_count"] is not None]
    actual_sum = sum(item["actual_count"] for item in comparable)
    planned_comparison = sum(item["planned_count"] for item in comparable)
    errors = [item["plan_error"] for item in comparable]
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("요약")
    summary.append(["항목", "값"])
    summary_rows = [
        ("기간", f"{start.isoformat()} ~ {end.isoformat()}"), ("식사유형", meal_type), ("지점", station_name or ""),
        ("강수필터", rain), ("메뉴필터", menu or ""), ("기온구간", temp_bucket or ""), ("서비스수", len(items)),
        ("계획합계", sum(item["planned_count"] for item in items)), ("비교 N", len(comparable)),
        ("비교 계획합계", planned_comparison), ("실제합계", actual_sum),
        ("WAPE(%)", _round(sum(abs(value) for value in errors) / actual_sum * 100) if actual_sum else None),
        ("Bias율(%)", _round(sum(errors) / actual_sum * 100) if actual_sum else None),
    ]
    for row in summary_rows:
        summary.append(row)
    detail = workbook.create_sheet("상세자료")
    detail.append(["일자", "식사유형", "계획식수", "실제식수", "계획-실제", "대표메뉴", "평균기온", "강수량", "평균습도", "적설량", "일조시간", "지점", "비고"])
    for item in items:
        weather = item["weather"] or {}
        detail.append([
            item["service_date"], item["meal_type"], item["planned_count"], item["actual_count"], item["plan_error"],
            ", ".join(item["representative_menus"]), weather.get("avg_temp"), weather.get("precipitation"),
            weather.get("avg_humidity"), weather.get("snow_depth"), weather.get("sunshine_hours"), resolved_id, item["note"],
        ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
